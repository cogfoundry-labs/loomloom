"""Offline tests for scripts/preflight.py and scripts/experiment.py (no gateway calls)."""
import base64
import copy
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "scripts"))
import experiment as ex  # noqa: E402
import ledger as lg  # noqa: E402
import preflight as pf  # noqa: E402
import workbook as wbk  # noqa: E402

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
PLAN = json.loads((HERE / "fixtures" / "plan-mighty.json").read_text(encoding="utf-8"))
NB = "google/gemini-2.5-flash-image"
SB = "openai/gpt-image-2.5-sunburst"
SEED = "bytedance/doubao-seedream-5-0-pro"

pf.il.load_price_observations = lambda: {}   # tests must not depend on a local observed-price cache
pf.REF_SUPPORT_FILE = HERE / "fixtures" / "reference-support.json"   # ...nor on the live reference-support data
ADVISOR = pf.Advisor(PLAN["intent"])        # real catalog + Arena data, built once


def make_experiment(tmp: Path, plan: dict, target=30, with_ref=True) -> Path:
    src = tmp / "src"
    (src / "refs").mkdir(parents=True)
    if with_ref:
        (src / "refs" / "mighty-product.png").write_bytes(PNG)
    (src / "plan.json").write_text(json.dumps(plan), encoding="utf-8")
    out = tmp / "exp"
    ex.build_experiment(src / "plan.json", out, target=target)
    return out


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def build(self, plan=None, **kw):
        self.exp = make_experiment(self.tmp, plan or PLAN, **kw)
        return self.exp

    def ledger(self):
        return lg.load(self.exp / "ledger.json")

    def edit_ledger(self, fn):
        led = self.ledger()
        fn(led)
        lg.save(self.exp / "ledger.json", led)
        (self.exp / "experiment.xlsx").unlink()          # ledger is the source in these tests

    def run_pf(self, **kw):
        return pf.run_preflight(self.exp, advisor=ADVISOR, **kw)


class Plan(Base):
    def test_plan_writes_ledger_workbook_and_ticks_the_recommended_rows(self):
        self.build()
        led = self.ledger()
        self.assertEqual(len(led["rows"]), 9)
        self.assertTrue(all(r["selected"] for r in led["rows"]))
        self.assertTrue((self.exp / "experiment.xlsx").exists())
        self.assertTrue((self.exp / "refs" / "mighty-product.png").exists())   # copied next to the experiment
        self.assertEqual({r["status"] for r in led["rows"]}, {"Draft"})

    def test_target_below_the_valid_count_ticks_only_the_batch(self):
        plan = copy.deepcopy(PLAN)
        plan["dimensions"]["mood"] = ["calm", "bold", "warm"]
        self.build(plan, target=12)
        led = self.ledger()
        self.assertEqual(len(led["rows"]), 27)                     # all valid combinations are in the workbook
        self.assertEqual(sum(r["selected"] for r in led["rows"]), 12)

    def test_spread_makes_model_an_explicit_dimension(self):
        plan = copy.deepcopy(PLAN)
        plan["model_strategy"] = "spread"
        self.build(plan, with_ref=False)
        led = self.ledger()
        self.assertEqual(led["experiment"]["dimensions"], ["camera", "lighting"])
        models = {r["model"] for r in led["rows"]}
        self.assertEqual(len(models), 2)
        for m in models:
            cams = {r["params"]["camera"] for r in led["rows"] if r["model"] == m}
            self.assertEqual(cams, {"eye-level front", "three-quarter high", "low angle"})   # each model covered


class ApprovalScope(Base):
    def test_preflight_says_what_the_fingerprint_authorizes(self):
        self.build()
        text = pf.format_report(self.run_pf())
        self.assertIn("authorizes exactly 9 images and nothing else", text)

    def test_resaving_an_unchanged_workbook_is_not_a_difference_whatever_its_file_time_says(self):
        import os
        self.build()
        fp = self.run_pf()["fingerprint"]
        wb = self.exp / "experiment.xlsx"
        later = wb.stat().st_mtime + 3600
        os.utime(wb, (later, later))                                              # "saved" an hour later with the same content
        self.assertEqual(ex.snapshot_drift(self.exp, fp), [])

    def test_a_changed_value_is_reported_by_row_and_field_even_if_the_file_time_did_not_move(self):
        import os
        self.build()
        fp = self.run_pf()["fingerprint"]
        led = self.ledger()
        first = led["rows"][0]
        dim = next(iter(first["params"]))
        other = next(r["params"][dim] for r in led["rows"] if r["params"][dim] != first["params"][dim])
        first["params"][dim] = other
        lg.save(self.exp / "ledger.json", led)
        wb = self.exp / "experiment.xlsx"
        wb.unlink()                                                               # the ledger is the source in this test
        drift = ex.snapshot_drift(self.exp, fp)
        self.assertTrue(any(line.startswith(first["id"] + ":") and "prompt" in line for line in drift), drift)
        self.assertEqual(len(drift), 1)                                           # only the edited row is named

    def test_a_row_unticked_after_approval_is_named_because_it_will_still_run(self):
        self.build()
        fp = self.run_pf()["fingerprint"]
        rid = self.ledger()["rows"][0]["id"]

        def untick(led):
            for r in led["rows"]:
                if r["id"] == rid:
                    r["selected"] = False
        self.edit_ledger(untick)
        drift = ex.snapshot_drift(self.exp, fp)
        self.assertEqual(len(drift), 1)
        self.assertIn("will still run", drift[0])

    def test_rows_that_already_ran_and_retry_snapshots_are_not_compared_and_unknown_fingerprints_are_quiet(self):
        self.build()
        fp = self.run_pf()["fingerprint"]
        def already_ran_and_edited(led):
            led["rows"][0]["status"] = "Completed"
            led["rows"][0]["params"][next(iter(led["rows"][0]["params"]))] = "zzz"
        self.edit_ledger(already_ran_and_edited)
        self.assertEqual(ex.snapshot_drift(self.exp, fp), [])                     # a resume: the row already ran
        self.assertEqual(ex.snapshot_drift(self.exp, "nosuchfingerprint"), [])

    def test_the_visual_checks_are_printed_after_a_run_as_unverified_and_absent_when_the_plan_has_none(self):
        checks = ["the text reads exactly what was asked", "no extra text anywhere"]
        self.build(dict(copy.deepcopy(PLAN), visual_checks=checks))
        note = ex.visual_checks_note(self.exp)
        self.assertIn("the run verifies none of these", note)
        self.assertIn("a model's reading or the user's own look", note)
        for i, check in enumerate(checks, 1):
            self.assertIn(f"{i}. {check}", note)
        plan = json.loads((self.exp / "plan.json").read_text(encoding="utf-8"))
        plan.pop("visual_checks")
        (self.exp / "plan.json").write_text(json.dumps(plan), encoding="utf-8")
        self.assertIsNone(ex.visual_checks_note(self.exp))
        self.assertIsNone(ex.visual_checks_note(self.tmp / "nowhere"))

    def test_ticked_rows_outside_the_approved_snapshot_are_listed_as_not_approved(self):
        self.build()
        rows = [r["id"] for r in self.ledger()["rows"]]
        rep = self.run_pf(only=rows[:2])                                          # a calibration-sized subset: 2 of 9 rows
        left = ex.unapproved_ticked(self.exp, rep["fingerprint"])
        self.assertEqual(sorted(left), sorted(rows[2:]))
        self.assertEqual(ex.unapproved_ticked(self.exp, self.run_pf()["fingerprint"]), [])       # the full preflight covers all nine
        self.assertEqual(ex.unapproved_ticked(self.exp, "nosuchfingerprint"), [])


class EndToEndFindings(Base):
    """Findings of the end-to-end test of the latest workflow."""
    def test_the_singular_is_used_for_one_image(self):
        self.build()
        rep = self.run_pf(only=[self.ledger()["rows"][0]["id"]])
        self.assertIn("authorizes exactly 1 image and nothing else", pf.format_report(rep))

    def test_a_plan_changed_after_the_build_is_reported_and_an_unchanged_one_is_not(self):
        self.build()
        self.assertFalse([i for i in self.run_pf()["info"] if "plan.json was changed" in i])
        plan = json.loads((self.exp / "plan.json").read_text(encoding="utf-8"))
        plan["visual_checks"] = ["a changed check"]
        (self.exp / "plan.json").write_text(json.dumps(plan), encoding="utf-8")
        info = " ".join(self.run_pf()["info"])
        self.assertIn("plan.json was changed after the experiment was built", info)
        self.assertIn("keep it", info)

    def test_acknowledging_a_person_notice_is_not_a_plan_change(self):
        self.build()
        plan = json.loads((self.exp / "plan.json").read_text(encoding="utf-8"))
        plan["consent_acknowledged"] = "2026-10-10T12:00:00"
        (self.exp / "plan.json").write_text(json.dumps(plan), encoding="utf-8")
        self.assertFalse([i for i in self.run_pf()["info"] if "plan.json was changed" in i])

    def test_an_old_ledger_without_a_plan_hash_says_nothing(self):
        self.build()

        def strip(led):
            led["experiment"].pop("plan_hash", None)
        self.edit_ledger(strip)
        self.assertFalse([i for i in self.run_pf()["info"] if "plan.json was changed" in i])

    def test_the_dry_run_gives_an_indicative_price_when_no_price_is_verified(self):
        import experiment as exp_mod
        plan = copy.deepcopy(PLAN)
        plan["quality"] = "medium"
        plan["references"] = []
        s = exp_mod.preview_plan(plan, 9)
        if s["estimate"]["unverified_rows"]:
            text = exp_mod.format_plan_summary(s)
            self.assertIn("Indicative:", text)
            self.assertIn("a hint, not a quote", text)
            self.assertIn("$0.0000 known", text)                     # the hint is never folded into the known total


class IndicativePrice(unittest.TestCase):
    HINTS = {"measured": __import__("datetime").date.today().isoformat(), "usd_per_image": {"m|1024x1536|q=medium": 0.01, "m|1152x1440|q=medium": 0.02, "m|1024x1536|q=high": 0.09}}

    def test_an_exact_hint_gives_a_single_figure_and_a_safe_limit(self):
        r = pf.indicative(["m|1024x1536|q=medium"] * 4, self.HINTS)
        self.assertEqual((r["low"], r["high"], r["rows"]), (0.04, 0.04, 4))
        self.assertEqual(r["max_usd"], 0.06)                            # 1.5 x 0.04, never below 0.05

    def test_an_unmatched_size_gives_the_range_for_that_model_and_quality_only(self):
        r = pf.indicative(["m|default|q=medium"] * 2, self.HINTS)
        self.assertEqual((r["low"], r["high"]), (0.02, 0.04))            # not the q=high price
        self.assertEqual(r["max_usd"], 0.06)

    def test_reference_and_auto_quality_rows_get_no_range(self):
        self.assertIsNone(pf.indicative(["m|default|q=medium|reference"], self.HINTS))        # a reference call bills about 3.4x
        self.assertIsNone(pf.indicative(["m|default"], self.HINTS))                           # auto quality bills differently
        self.assertIsNone(pf.indicative(["m|1024x1536|q=medium", "m|default|reference"], self.HINTS))   # one unhintable row: no total
        exact = dict(self.HINTS, usd_per_image={**self.HINTS["usd_per_image"], "m|default|q=medium|reference": 0.05})
        self.assertEqual(pf.indicative(["m|default|q=medium|reference"], exact)["high"], 0.05)   # an exact reference hint is fine

    def test_old_or_undated_hints_are_not_shown(self):
        old = dict(self.HINTS, measured="2020-01-01")
        self.assertIsNone(pf.indicative(["m|1024x1536|q=medium"], old))
        self.assertIsNone(pf.indicative(["m|1024x1536|q=medium"], dict(self.HINTS, measured="soon")))

    def test_the_price_key_matches_the_observed_price_keys(self):
        self.assertEqual(pf.price_key("m", "1024x1536", "medium", "text"), "m|1024x1536|q=medium")
        self.assertEqual(pf.price_key("m", None, "auto", "text"), "m|default")
        self.assertEqual(pf.price_key("m", None, None, "reference"), "m|default|reference")
        self.assertEqual(pf.price_key("m", "1K", "high", "reference"), "m|1K|q=high|reference")

    def test_no_matching_model_or_no_hints_gives_nothing(self):
        self.assertIsNone(pf.indicative(["other|default|q=medium"], self.HINTS))
        self.assertIsNone(pf.indicative(["m|default"], {}))

    def test_the_report_prints_the_hint_but_keeps_it_out_of_the_known_total(self):
        rep = {"ready": 4, "selected": 32, "issues": [], "images": 4, "models": ["m"], "references": [], "by_model": {"m": 4}, "known_usd": 0.0,
               "unverified_rows": ["r1"] * 4, "warnings": [], "info": [], "merge_warnings": [], "workbook_saved": None, "fingerprint": "abc",
               "coverage": 0.99, "subset": True, "indicative": pf.indicative(["m|1024x1536|q=medium"] * 4, self.HINTS), "retry": False}
        text = pf.format_report(rep)
        self.assertIn("$0.0000 known", text)
        self.assertIn("Indicative:         about $0.04", text)
        self.assertIn("--max-usd 0.06", text)
        self.assertIn("this batch is a subset", text)
        self.assertIn("all ticked rows, not only this subset", text)

    def test_the_shipped_hints_file_loads(self):
        self.assertTrue(pf.load_price_hints()["usd_per_image"])

    def test_the_target_note_says_what_happened_to_the_target(self):
        r = {"by": "direction"}
        self.assertIn("rounded up", ex._target_note({"target": 30, "recommended": 32}, r))
        self.assertIn("only 20 images exist for a target of 30", ex._target_note({"target": 30, "recommended": 20}, r))
        self.assertEqual(ex._target_note({"target": 30, "recommended": 30}, r), "your target of 30: ")
        self.assertEqual(ex._target_note({"target": 0, "recommended": 12}, r), "")
        self.assertEqual(ex._target_note({"recommended": 12}, r), "")


class Preflight(Base):
    def test_ready_rows_snapshot_and_fingerprint(self):
        self.build()
        rep = self.run_pf()
        self.assertEqual((rep["ready"], rep["selected"], len(rep["issues"])), (9, 9, 0))
        fp = rep["fingerprint"]
        snap = json.loads((self.exp / "snapshots" / f"{fp}.json").read_text(encoding="utf-8"))
        self.assertEqual(snap["fingerprint"], fp)
        row = snap["rows"][0]
        self.assertEqual(row["mode"], "reference")
        self.assertEqual(row["references"][0]["sha256"], pf.sha256_file(self.exp / "refs" / "mighty-product.png"))
        self.assertEqual(pf.snapshot_fingerprint(snap), fp)       # the fingerprint covers the whole snapshot body
        self.assertEqual(self.run_pf()["fingerprint"], fp)          # deterministic

    def test_one_per_dimension_prices_a_calibration_subset_and_leaves_the_rest_ticked(self):
        self.build()
        rep = self.run_pf(one_per="camera")
        self.assertEqual((rep["ready"], len(rep["issues"])), (3, 0))          # three cameras, one row each
        led = self.ledger()
        self.assertTrue(all(r["selected"] for r in led["rows"]))               # the other ticked rows stay ticked
        self.assertEqual(sum(r["status"] == "Ready" for r in led["rows"]), 3)
        self.assertEqual(self.run_pf(one_per="camera")["fingerprint"], rep["fingerprint"])
        self.assertNotEqual(self.run_pf()["fingerprint"], rep["fingerprint"])  # the plain preflight takes all nine

    def test_only_takes_named_rows_and_refuses_unknown_names(self):
        self.build()
        first = self.ledger()["rows"][0]["id"]
        rep = self.run_pf(only=[first])
        self.assertEqual(rep["ready"], 1)
        bad = self.run_pf(only=["r999"], one_per="colour")
        problems = " ".join(p for i in bad["issues"] for p in i["problems"])
        self.assertIn("no row r999", problems)
        self.assertIn("no dimension 'colour'", problems)

    def test_reference_rows_prefer_the_verified_model_and_it_is_priced(self):
        self.build()
        rep = self.run_pf()
        self.assertEqual(rep["models"], [NB])
        self.assertAlmostEqual(rep["known_usd"], 9 * 0.039, places=6)
        self.assertEqual(rep["unverified_rows"], [])

    def test_documented_reference_model_is_allowed_but_unverified_and_unpriced(self):
        self.build()
        self.edit_ledger(lambda led: [r.update(model=SB) for r in led["rows"]])
        rep = self.run_pf()
        self.assertEqual(rep["ready"], 9)
        self.assertEqual(len(rep["unverified_rows"]), 9)
        self.assertEqual(rep["known_usd"], 0)                       # never a made-up number
        self.assertTrue(any("documented" in w for w in rep["warnings"]))
        self.assertIn("price unverified", rep["text"])

    def test_text_mode_prices_from_the_catalog_one_image_per_row(self):
        text_plan = copy.deepcopy(PLAN)
        text_plan["references"] = []
        self.build(text_plan, with_ref=False)
        self.edit_ledger(lambda led: [r.update(model=SB) for r in led["rows"]])
        rep = self.run_pf()
        r1 = next(r for r in rep["snapshot"]["rows"] if r["id"] == "r001")
        self.assertEqual((r1["mode"], r1["model"], r1["qty"]), ("text", SB, 1))
        self.assertAlmostEqual(r1["est_usd"], 0.0069, places=5)
        self.assertEqual(rep["images"], rep["ready"])                  # images = ticked rows

    def test_invalid_rows_become_draft_and_are_not_in_the_snapshot(self):
        self.build()

        def bad(led):
            led["rows"][0]["take"] = 0
            led["rows"][1]["params"]["lighting"] = ""
            led["rows"][2]["reference"] = "nope"
        self.edit_ledger(bad)
        rep = self.run_pf()
        reasons = {i["id"]: " ".join(i["problems"]) for i in rep["issues"]}
        self.assertIn("Take must be a whole number from 1", reasons["r001"])
        self.assertIn("missing value for lighting", reasons["r002"])
        self.assertIn("unknown reference", reasons["r003"])
        self.assertEqual(rep["ready"], 6)
        ids = {r["id"] for r in rep["snapshot"]["rows"]}
        self.assertTrue({"r001", "r002", "r003"}.isdisjoint(ids))
        self.assertEqual({r["id"]: r["status"] for r in self.ledger()["rows"]}["r001"], "Draft")

    def test_missing_or_unreadable_reference_file(self):
        self.build()
        (self.exp / "refs" / "mighty-product.png").write_bytes(b"not an image")
        self.assertIn("not a PNG, JPEG or WebP image", self.run_pf()["issues"][0]["problems"][0])
        (self.exp / "refs" / "mighty-product.png").unlink()
        self.assertIn("not found", self.run_pf()["issues"][0]["problems"][0])

    def test_unselected_removed_and_unknown_rows_are_not_included(self):
        self.build()

        def mark(led):
            led["rows"][0]["selected"] = False
            led["rows"][1]["status"] = "Removed"
            led["rows"][2]["status"] = "Unknown"
        self.edit_ledger(mark)
        rep = self.run_pf()
        self.assertEqual(rep["ready"], 6)
        self.assertTrue(any("r003" in w and "never retried automatically" in w for w in rep["warnings"]))

    def test_unreliable_value_is_reported_in_reference_mode_only(self):
        self.build()
        self.assertTrue(any('"low angle" is unreliable' in i for i in self.run_pf()["info"]))
        text_plan = copy.deepcopy(PLAN)
        text_plan["references"] = []
        shutil.rmtree(self.tmp / "exp")
        shutil.rmtree(self.tmp / "src")
        self.exp = make_experiment(self.tmp, text_plan, with_ref=False)
        self.assertFalse(any("unreliable" in i for i in self.run_pf()["info"]))

    def test_custom_value_warning_and_generic_wording_in_the_snapshot(self):
        self.build()

        def custom(led):
            led["rows"][0]["params"]["lighting"] = "rainy night"
        self.edit_ledger(custom)
        rep = self.run_pf()
        self.assertTrue(any('custom value "rainy night"' in w for w in rep["warnings"]))
        r1 = next(r for r in rep["snapshot"]["rows"] if r["id"] == "r001")
        self.assertIn("rainy night as the lighting", r1["prompt"])

    def test_changing_a_value_changes_the_fingerprint(self):
        self.build()
        before = self.run_pf()["fingerprint"]

        def change(led):
            led["rows"][0]["params"]["lighting"] = "golden hour"
            led["rows"][0]["params"]["camera"] = "eye-level front"
        self.edit_ledger(change)
        self.assertNotEqual(self.run_pf()["fingerprint"], before)

    def test_aspect_ratio_that_the_model_cannot_honor_is_reported(self):
        plan = copy.deepcopy(PLAN)
        plan["references"] = []
        plan["dimensions"]["aspect"] = ["4:5"]
        self.build(plan, with_ref=False, target=9)
        self.edit_ledger(lambda led: [r.update(model=SEED) for r in led["rows"]])
        rep = self.run_pf()
        self.assertTrue(any("requested 4:5 ->" in i for i in rep["info"]))
        row = rep["snapshot"]["rows"][0]
        self.assertTrue(row["size"])                                  # a real size token was chosen

    def test_unknown_model_override_is_an_issue(self):
        self.build()
        self.edit_ledger(lambda led: led["rows"][0].update(model="nope/model"))
        self.assertIn("not in the catalog", self.run_pf()["issues"][0]["problems"][0])

    def test_estimate_only_mode_touches_no_file(self):
        self.build()
        before = sorted(p.name for p in self.exp.rglob("*") if p.is_file())
        rep = pf.run_preflight(self.exp, advisor=ADVISOR, write=False)
        self.assertEqual(sorted(p.name for p in self.exp.rglob("*") if p.is_file()), before)
        self.assertEqual(rep["ready"], 9)


class EditedWorkbook(Base):
    def test_user_edits_in_excel_drive_the_snapshot(self):
        self.build()
        import openpyxl
        path = self.exp / "experiment.xlsx"
        book = openpyxl.load_workbook(path)
        ws = book["Experiment"]
        ws["B2"] = False                      # untick r001
        ws["F3"] = 2                          # r002 Take 2
        ws["C4"] = "low angle"                # r003 camera
        book.save(path)
        rep = pf.run_preflight(self.exp, advisor=ADVISOR)
        self.assertEqual(rep["ready"], 8)
        rows = {r["id"]: r for r in rep["snapshot"]["rows"]}
        self.assertNotIn("r001", rows)
        self.assertEqual(rep["snapshot"]["rows"][0]["qty"], 1)         # one image per row, whatever Take says
        self.assertEqual(self.ledger()["rows"][1]["take"], 2)
        self.assertIn("Low camera position at counter level", rows["r003"]["prompt"])   # reference wording
        self.assertEqual(rep["images"], 8)                                              # one image per ticked row
        self.assertTrue(rep["workbook_saved"])


class Check(unittest.TestCase):
    def test_data_files_are_consistent(self):
        self.assertEqual(ex.check(HERE / "fixtures" / "plan-mighty.json"), [])

    def test_bad_plan_is_reported(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            (tmp / "p.json").write_text(json.dumps({"schema_version": 1}), encoding="utf-8")
            self.assertTrue(any("plan.json is invalid" in p for p in ex.check(tmp / "p.json")))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class Grouping(unittest.TestCase):
    def test_group_messages(self):
        msgs = ["r001: x", "r002: x", "plain", "r003: y"] + [f"r{n:03d}: z" for n in range(10, 18)]
        out = pf.group_messages(msgs)
        self.assertIn("plain", out)
        self.assertIn("x: r001, r002", out)
        self.assertIn("y: r003", out)
        self.assertTrue(any(o.startswith("z: r010, r011, r012, r013, r014, ... (8 rows)") for o in out))


class Refresh(Base):
    def test_refresh_merges_edits_regenerates_and_keeps_prev(self):
        self.build()
        import openpyxl
        path = self.exp / "experiment.xlsx"
        book = openpyxl.load_workbook(path)
        book["Experiment"]["F2"] = 3
        book.save(path)
        res = ex.refresh(self.exp)
        self.assertTrue((self.exp / "experiment.prev.xlsx").exists())
        again = wbk.read_workbook(self.exp / "experiment.xlsx", ["camera", "lighting"])
        self.assertEqual(again["rows"][0]["take"], 3)
        self.assertEqual(res["warnings"], [])
        self.assertEqual(self.ledger()["rows"][0]["take"], 3)


if __name__ == "__main__":
    unittest.main()
