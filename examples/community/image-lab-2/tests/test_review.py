"""Regression tests for the findings of the code review against design-v2.md (2026-10-07).

One test (or a few) per verified finding. Names say what used to go wrong."""
import contextlib
import copy
import http.server
import io
import json
import os
import subprocess
import sys
import threading
import time
import unittest
from pathlib import Path
from unittest import mock

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(HERE))
import compile as cp  # noqa: E402
import experiment as ex  # noqa: E402
import generate as gen  # noqa: E402
import image as il  # noqa: E402
import ledger as lg  # noqa: E402
import matrix as mx  # noqa: E402
import preflight as pf  # noqa: E402
import sheet as sh  # noqa: E402
import test_generate as tg  # noqa: E402
import workbook as wbk  # noqa: E402
from test_generate import FakeGateway, quiet  # noqa: E402
from test_m3 import ledger_select  # noqa: E402
from test_preflight import ADVISOR, Base, NB, PLAN, PNG, make_experiment  # noqa: E402

ENV = {**os.environ, "PYTHONUTF8": "1"}


class RunBase(Base):
    prepare = tg.GenerateTests.prepare
    run_gen = tg.GenerateTests.run_gen

    def retry_pf(self, **kw):
        return pf.run_preflight(self.exp, advisor=ADVISOR, retry=True, **kw)


# --------------------------------------------------------------------------- #
# retry must not pay twice
# --------------------------------------------------------------------------- #
class RetryAccountingTests(RunBase):
    def test_an_unknown_sample_is_not_resubmitted(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway(lambda n, b: (0, {}, "TimeoutError") if n == 1 else None), concurrency=1)
        self.assertEqual(self.ledger()["rows"][0]["status"], "Unknown")
        rep = self.retry_pf()
        self.assertEqual(rep["ready"], 0)                                   # hard rule: Unknown is never auto-retried
        self.assertTrue(any("unknown" in w.lower() for w in rep["warnings"]))
        self.assertEqual(self.retry_pf(include_unknown=True)["ready"], 1)   # only when the user asks

    def test_a_blocked_sample_is_not_retried(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway(fail_reasons={"req-1": "blocked by safety policy"}), concurrency=1)
        self.assertEqual(self.ledger()["rows"][0]["status"], "Blocked")
        rep = self.retry_pf()
        self.assertEqual(rep["ready"], 0)
        self.assertIn("Blocked", rep["text"])                              # explained: the model refused, not retried

    def test_retry_does_not_duplicate_samples_that_a_paused_batch_still_holds(self):
        fp = self.prepare(3)
        res = self.run_gen(fp, FakeGateway(), concurrency=1, max_usd=0.05)     # 1 completes, 2 stay Pending
        self.assertEqual(res["by_status"], {"Completed": 1, "Pending": 2})
        rep = self.retry_pf()
        self.assertEqual(rep["ready"], 0)                                      # the held rows are not Failed: nothing to retry
        res2 = self.run_gen(fp, FakeGateway())                                 # resuming finishes them, no duplicates
        self.assertEqual(res2["by_status"], {"Completed": 3})

    def test_images_billed_but_not_downloaded_are_recovered_not_regenerated(self):
        fp = self.prepare(1)

        class Broken(FakeGateway):
            def download(self, url, dest):
                raise OSError("tunnel connection failed")

        self.run_gen(fp, Broken())
        a = self.ledger()["attempts"][0]
        self.assertEqual((a["status"], a["cost"] > 0, bool(a["url"])), ("Failed", True, True))
        rep = self.retry_pf()
        self.assertEqual(rep["ready"], 0)                                      # no second bill
        self.assertTrue(any("recover" in w for w in rep["warnings"]))
        res = gen.recover_downloads(self.exp, gateway=FakeGateway(), log=quiet, record_prices=False)
        self.assertEqual(res["recovered"], ["r001"])

    def test_sample_ids_are_never_reused_across_batches(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway())
        self.run_gen(fp, FakeGateway(), again=True)                            # the same snapshot, run again
        led = self.ledger()
        led["rows"][0]["params"]["lighting"] = "dramatic spotlight"             # the row is edited: a new configuration
        lg.save(self.exp / "ledger.json", led)
        rep = pf.run_preflight(self.exp, advisor=ADVISOR)                      # and a fresh preflight of the ticked row
        self.run_gen(rep["fingerprint"], FakeGateway())
        ids = [a["sample_id"] for a in self.ledger()["attempts"]]
        self.assertEqual(len(ids), len(set(ids)), ids)
        self.assertEqual(ids, ["r001", "r001-2", "r001-3"])

    def test_preflight_leaves_the_history_of_rows_that_already_ran(self):
        fp = self.prepare(2)
        self.run_gen(fp, FakeGateway(fail_reasons={"req-1": "boom"}), concurrency=1)
        before = {r["id"]: r["status"] for r in self.ledger()["rows"] if r["selected"]}
        self.assertEqual(sorted(before.values()), ["Completed", "Failed"])
        pf.run_preflight(self.exp, advisor=ADVISOR)                            # an ordinary preflight, then decline
        after = {r["id"]: r["status"] for r in self.ledger()["rows"] if r["selected"]}
        self.assertEqual(after, before)
        self.assertEqual(self.retry_pf()["ready"], 1)                          # so retry can still find the Failed row

    def test_an_unfinished_batch_is_pointed_out_with_its_fingerprint(self):
        fp = self.prepare(2)
        self.run_gen(fp, FakeGateway(polls=10_000), poll_timeout=0.05)
        rep = pf.run_preflight(self.exp, advisor=ADVISOR)
        self.assertTrue(any("unfinished" in w and fp in w for w in rep["warnings"]), rep["warnings"])


# --------------------------------------------------------------------------- #
# one writer at a time
# --------------------------------------------------------------------------- #
class LockTests(RunBase):
    def test_writers_are_refused_while_a_run_holds_the_ledger(self):
        fp = self.prepare(1)
        with lg.Lock(self.exp, "run test"):
            with self.assertRaises(lg.LedgerBusy):
                pf.run_preflight(self.exp, advisor=ADVISOR)
            with self.assertRaises(lg.LedgerBusy):
                ex.refresh(self.exp)
            with self.assertRaises(lg.LedgerBusy):
                gen.recover_downloads(self.exp, gateway=FakeGateway(), log=quiet)
            with self.assertRaises(lg.LedgerBusy):
                ex.build_quick("a cat", "e-commerce product photo", 1, self.exp)
            with self.assertRaises(lg.LedgerBusy):                              # a second run of the same experiment
                self.run_gen(fp, FakeGateway())
            self.assertEqual(pf.run_preflight(self.exp, advisor=ADVISOR, write=False)["ready"], 1)   # reads are fine
        self.assertEqual(pf.run_preflight(self.exp, advisor=ADVISOR)["ready"], 1)                   # released: works again

    def test_the_lock_is_released_after_a_run_and_a_stale_one_is_taken_over(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway())
        self.assertFalse((self.exp / lg.LOCK_NAME).exists())
        (self.exp / lg.LOCK_NAME).write_text(json.dumps({"pid": 1, "who": "crashed", "started": "x",
                                                          "beat": time.time() - 3600}), encoding="utf-8")
        lg.assert_unlocked(self.exp)                                            # stale: not busy
        self.run_gen(fp, FakeGateway(), again=True)                             # and `run` takes it over
        self.assertFalse((self.exp / lg.LOCK_NAME).exists())

    def test_the_cli_reports_busy_and_exits_2_without_a_traceback(self):
        self.prepare(1)
        with lg.Lock(self.exp, "run test"):
            r = subprocess.run([sys.executable, str(SCRIPTS / "image.py"), "preflight", "--dir", str(self.exp)],
                               capture_output=True, text=True, encoding="utf-8", env=ENV)
        self.assertEqual(r.returncode, 2)
        self.assertIn("BUSY", r.stderr)
        self.assertNotIn("Traceback", r.stderr)


# --------------------------------------------------------------------------- #
# the snapshot, the budget, the gateway
# --------------------------------------------------------------------------- #
    def test_a_row_edited_after_it_ran_does_not_inherit_the_old_image(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway())                                                  # r001 is Completed
        led = self.ledger()
        old_params = dict(led["rows"][0]["params"])
        self.assertEqual(led["attempts"][0]["params"], old_params)                       # recorded when the image was made
        led["rows"][0]["params"]["lighting"] = "dramatic spotlight"                       # the user edits the row
        lg.save(self.exp / "ledger.json", led)
        rep = pf.run_preflight(self.exp, advisor=ADVISOR)                                # a new configuration: it is generated again
        self.assertEqual(rep["ready"], 1)
        self.assertTrue(any("edited after it was generated" in i for i in rep["info"]))
        self.assertEqual(rep["snapshot"]["rows"][0]["first_sample"], 2)                  # r001-2: ids are never reused
        self.run_gen(rep["fingerprint"], FakeGateway())
        led = self.ledger()
        by_sample = {a["sample_id"]: a["params"]["lighting"] for a in led["attempts"] if a["status"] == "Completed"}
        self.assertEqual(by_sample["r001"], old_params["lighting"])                      # the old image keeps its own label
        self.assertEqual(by_sample["r001-2"], "dramatic spotlight")
        d = sh.build_data(self.exp)
        labels = {s["id"]: s["params"]["lighting"] for r in d["rows"] for s in r["samples"] if s["status"] == "Completed"}
        self.assertEqual(labels["r001"], old_params["lighting"])                         # and so does the contact sheet
        import openpyxl
        wbk_path = self.exp / "experiment.xlsx"
        import experiment as ex2
        ex2.refresh(self.exp)
        ws = openpyxl.load_workbook(wbk_path)["Experiment"]
        head = [c.value for c in ws[1]]
        self.assertEqual(ws.cell(row=2, column=head.index("File") + 1).value, "r001-2.png")   # the row shows its latest image

    def test_a_finished_batch_rejected_before_any_billing_can_be_run_again_without_again(self):
        fp = self.prepare(2)
        bad = FakeGateway(lambda n, b: (400, {"error": {"code": "bad_request", "message": "no"}}, ""))   # no breaker
        res = self.run_gen(fp, bad)
        self.assertEqual(res["by_status"], {"Failed": 2})
        self.assertFalse(gen.batch_spent(self.ledger(), 1))
        logs = []
        res2 = gen.run_batch(self.exp, fp, gateway=FakeGateway(), poll_seconds=0.01, record_prices=False, log=logs.append)
        self.assertEqual(res2["by_status"], {"Completed": 2})                      # no --again, no false 'new spend' refusal
        self.assertTrue(any("rejected before anything was billed" in m for m in logs))
        with self.assertRaises(gen.RunRefused):                                    # but a batch that spent is still protected
            self.run_gen(fp, FakeGateway())
        led = self.ledger()
        self.assertEqual(sorted(a["sample_id"] for a in led["attempts"] if not a.get("superseded")), ["r001", "r002"])
        self.assertTrue(all(a.get("superseded") for a in led["attempts"] if a["batch"] == 1))   # kept, but holds no number
        d = sh.build_data(self.exp)
        self.assertEqual((d["stats"]["samples"], d["stats"]["problems"]), (2, 0))             # and is not shown on the sheet

    def test_a_wrong_key_leaves_samples_pending_and_the_fix_resubmits_all_of_them(self):
        fp = self.prepare(4)
        bad = FakeGateway(lambda n, b: (401, {"error": {"code": "invalid_api_key", "message": "bad key"}}, ""))
        res = self.run_gen(fp, bad, concurrency=1)                                  # the breaker stops after the first
        self.assertEqual(res["by_status"], {"Failed": 1, "Pending": 3})
        self.assertIn("rejected", res["stopped"])
        res2 = self.run_gen(fp, FakeGateway(), concurrency=1)                       # same fingerprint, the key is fixed
        self.assertEqual(res2["by_status"], {"Completed": 4})                       # including the one that was rejected
        self.assertEqual(sorted(a["sample_id"] for a in self.ledger()["attempts"]), ["r001", "r002", "r003", "r004"])

    def test_retry_with_nothing_to_retry_says_why_instead_of_printing_zeros(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway())
        rep = self.retry_pf()
        self.assertIn("Nothing to retry", rep["text"])
        self.assertNotIn("Pairwise coverage", rep["text"])
        self.assertNotIn("Balance", rep["text"])

    def test_a_batch_with_an_unknown_or_billed_sample_counts_as_spent(self):
        fp = self.prepare(2)
        self.run_gen(fp, FakeGateway(lambda n, b: (0, {}, "TimeoutError") if n == 1 else (401, {"error": {"code": "invalid_api_key"}}, "")),
                     concurrency=1)
        self.assertTrue(gen.batch_spent(self.ledger(), 1))

    def test_preflight_and_plan_summaries_show_the_model_and_exact_cents(self):
        fp = self.prepare(2)
        rep = pf.run_preflight(self.exp, advisor=ADVISOR)
        self.assertIn("image(s)", rep["text"])
        self.assertTrue(any(m in rep["text"] for m in rep["models"]))
        self.assertEqual(rep["by_model"], {rep["models"][0]: 2})
        s = ex.format_plan_summary(ex.build_experiment(self.tmp / "src" / "plan.json", self.tmp / "dry", dry_run=True))
        self.assertRegex(s, r"Estimated cost:\s+\$\d+\.\d{4} known")


class SnapshotAndBudgetTests(RunBase):
    def test_the_fingerprint_covers_the_numbers_the_spending_limit_comes_from(self):
        fp = self.prepare(2)
        p = self.exp / "snapshots" / f"{fp}.json"
        snap = json.loads(p.read_text(encoding="utf-8"))
        for key, val in (("estimated_usd_known", 99.0), ("unverified_rows", ["r001"])):
            tampered = copy.deepcopy(snap)
            tampered[key] = val
            self.assertNotEqual(pf.snapshot_fingerprint(tampered), fp, key)
        snap["estimated_usd_known"] = 99.0
        p.write_text(json.dumps(snap), encoding="utf-8")
        gw = FakeGateway()
        with self.assertRaises(gen.RunRefused):
            self.run_gen(fp, gw)
        self.assertEqual(gw.n, 0)

    def test_billed_failed_downloads_count_toward_the_spending_limit(self):
        fp = self.prepare(6)

        class Broken(FakeGateway):
            def download(self, url, dest):
                raise OSError("cdn outage")

        gw = Broken(cost=0.05)
        res = self.run_gen(fp, gw, concurrency=1, max_usd=0.12)
        self.assertEqual(gw.n, 2)                                              # not all 6: the failures were billed
        self.assertIn("budget guard", res["stopped"])

    def test_unpriced_samples_go_one_at_a_time_until_a_real_price_is_known(self):
        fp = self.prepare(4)
        p = self.exp / "snapshots" / f"{fp}.json"
        snap = json.loads(p.read_text(encoding="utf-8"))
        for r in snap["rows"]:
            r["est_usd"] = None
        snap["unverified_rows"] = [r["id"] for r in snap["rows"]]
        snap["fingerprint"] = fp2 = pf.snapshot_fingerprint(snap)
        (self.exp / "snapshots" / f"{fp2}.json").write_text(json.dumps(snap), encoding="utf-8")

        class Tracking(FakeGateway):
            def __init__(self):
                super().__init__(polls=3)
                self.inflight, self.at_submit = set(), []

            def submit(self, body):
                with self.lock:
                    self.at_submit.append(len(self.inflight))
                st, resp, err = super().submit(body)
                with self.lock:
                    self.inflight.add(resp["data"]["request_id"])
                return st, resp, err

            def poll(self, rid):
                st, resp, err = super().poll(rid)
                if resp["data"]["status"] == "COMPLETED":
                    with self.lock:
                        self.inflight.discard(rid)
                return st, resp, err

        gw = Tracking()
        res = self.run_gen(fp2, gw, concurrency=4, max_usd=5)
        self.assertEqual(res["by_status"], {"Completed": 4})
        self.assertEqual(gw.at_submit[1], 0)         # the 2nd sample waited until the 1st had a real price

    def test_the_guard_switches_to_real_prices_when_the_model_bills_more_than_estimated(self):
        fp = self.prepare(9)                                             # estimated at about $0.039 each
        gw = FakeGateway(cost=0.2, polls=1)                              # but really bills $0.20 each (5x)
        res = self.run_gen(fp, gw, concurrency=1, max_usd=0.50)
        self.assertLessEqual(gw.n, 3)                                    # 0.2 + 0.2 + next would exceed 0.5: stop at once
        self.assertLessEqual(res["actual_usd"], 0.50 + 1e-9)
        self.assertIn("budget guard", res["stopped"])

    def test_in_flight_samples_are_counted_at_the_real_price_not_the_estimate(self):
        fp = self.prepare(8)
        gw = FakeGateway(cost=0.2, polls=4)
        res = self.run_gen(fp, gw, concurrency=4, max_usd=0.9)           # 4 in flight at once, each really $0.20
        # the first batch of 4 is submitted at the estimate (no price known yet); after the first completes the guard
        # must count the rest at $0.20 and stop long before all 8 are sent
        self.assertLessEqual(res["actual_usd"], 1.3)
        self.assertGreater(res["by_status"].get("Pending", 0), 0)

    def test_the_observed_price_is_the_mean_of_the_batch_not_the_last_image(self):
        import json as _json
        fp = self.prepare(4)
        costs = iter([0.03, 0.055, 0.055, 0.03])

        class Uneven(FakeGateway):
            def poll(self, rid):
                st, resp, err = super().poll(rid)
                if resp["data"].get("status") == "COMPLETED":
                    resp["data"]["cost"] = [0.03, 0.055, 0.055, 0.03][int(rid.split("-")[1]) - 1]
                return st, resp, err

        obs_file = self.tmp / "price-observations.json"
        support_copy = self.tmp / "reference-support.json"
        import shutil as _sh
        _sh.copy(pf.REF_SUPPORT_FILE, support_copy)                       # promotion must never touch the frozen fixture
        old, old_support = gen.il.PRICE_OBSERVATIONS_FILE, pf.REF_SUPPORT_FILE
        gen.il.PRICE_OBSERVATIONS_FILE, pf.REF_SUPPORT_FILE = str(obs_file), support_copy
        try:
            gen.run_batch(self.exp, fp, gateway=Uneven(), poll_seconds=0.01, concurrency=1, max_usd=5, log=quiet)
        finally:
            gen.il.PRICE_OBSERVATIONS_FILE, pf.REF_SUPPORT_FILE = old, old_support
        obs = _json.loads(obs_file.read_text(encoding="utf-8"))
        self.assertEqual(len(obs), 1)
        self.assertAlmostEqual(next(iter(obs.values())), 0.04250, places=4)

    def test_a_cancelled_or_expired_task_is_a_failure_not_a_480_second_wait(self):
        fp = self.prepare(1)

        class Cancelled(FakeGateway):
            def poll(self, rid):
                return 200, {"data": {"status": "CANCELED"}}, ""

        t0 = time.time()
        res = self.run_gen(fp, Cancelled(), poll_timeout=30)
        self.assertLess(time.time() - t0, 5)
        self.assertEqual(res["by_status"], {"Failed": 1})

    def test_a_502_or_500_on_submit_is_unknown_never_retried(self):
        for code in (500, 502, 504):
            self.tearDown()
            self.setUp()
            fp = self.prepare(1)
            gw = FakeGateway(lambda n, b, c=code: (c, {}, f"HTTP {c}"))
            res = self.run_gen(fp, gw)
            self.assertEqual(res["by_status"], {"Unknown": 1}, code)
            self.assertEqual(gw.n, 1, code)

    def test_only_a_429_is_retried_on_a_submit_and_only_once(self):
        calls = []

        def fake_req(method, path, tok, body=None, retries=1):
            calls.append((method, retries))
            return responses.pop(0)

        g = gen.Gateway("tok")
        with mock.patch.object(il, "_req", fake_req), mock.patch.object(il, "RETRY_SLEEP_SECONDS", 0):
            responses = [(429, {}, "HTTP 429"), (200, {"data": {"request_id": "x"}}, "")]
            self.assertEqual(g.submit({})[0], 200)
            responses = [(502, {}, "HTTP 502")]
            calls.clear()
            self.assertEqual(g.submit({})[0], 502)
            self.assertEqual(len(calls), 1)
            responses = [(429, {}, ""), (429, {}, "")]
            calls.clear()
            self.assertEqual(g.submit({})[0], 429)
            self.assertEqual(len(calls), 2)
            self.assertTrue(all(r == 0 for _, r in calls))            # il._req itself never retries the POST

    def test_a_crashing_worker_does_not_abandon_the_batch(self):
        fp = self.prepare(3)

        class Boom(FakeGateway):
            def submit(self, body):
                if self.n == 1:
                    self.n += 1
                    raise RuntimeError("socket exploded")
                return super().submit(body)

        res = self.run_gen(fp, Boom(), concurrency=1)
        self.assertEqual(res["by_status"], {"Completed": 2, "Unknown": 1})     # the crashed one: may be billed
        led = self.ledger()
        self.assertIsNotNone(led["batches"][0]["ended_at"])                    # finish() ran
        self.assertNotIn("Queued", {r["status"] for r in led["rows"]})
        self.assertTrue((self.exp / "round-1" / "run.json").exists())

    def test_the_token_is_resolved_before_anything_is_written(self):
        fp = self.prepare(1)
        with mock.patch.object(gen.Gateway, "__init__", side_effect=SystemExit("no token")):
            with self.assertRaises(SystemExit):
                gen.run_batch(self.exp, fp, log=quiet)
        led = self.ledger()
        self.assertEqual((led["batches"], led["attempts"]), ([], []))           # no half-created batch
        self.assertFalse((self.exp / lg.LOCK_NAME).exists())
        self.assertNotIn("Queued", {r["status"] for r in led["rows"]})


class ReferenceSafetyTests(Base):
    def srow(self, sha, rel="refs/x.png"):
        return {"id": "r001", "model": NB, "prompt": "p", "size": None, "aspect_ratio": None,
                "references": [{"id": "ref1", "file": rel, "sha256": sha}]}

    def test_the_reference_is_re_hashed_at_send_time_and_confined(self):
        (self.tmp / "refs").mkdir()
        f = self.tmp / "refs" / "x.png"
        f.write_bytes(PNG)
        good = pf.sha256_file(f)
        cat = pf.il.load_model_catalog()
        self.assertIn("image", gen.build_request(self.srow(good), cat, self.tmp))
        f.write_bytes(PNG + b"changed")
        with self.assertRaises(gen.RunRefused):
            gen.build_request(self.srow(good), cat, self.tmp)                  # changed between approval and send
        with self.assertRaises(gen.RunRefused):
            gen.build_request(self.srow(good, "../outside.png"), cat, self.tmp / "refs")
        f.write_bytes(b"RIFF\x00\x00\x00\x00WAVEfmt ")
        with self.assertRaises(gen.RunRefused):
            gen.build_request(self.srow(pf.sha256_file(f)), cat, self.tmp)     # RIFF alone is not a WebP

    def test_oversized_or_foreign_references_are_issues_at_preflight(self):
        (self.tmp / "refs").mkdir()
        f = self.tmp / "refs" / "x.png"
        f.write_bytes(PNG)
        self.assertIsNone(pf.reference_problem(self.tmp, "refs/x.png"))
        with mock.patch.object(pf, "MAX_REFERENCE_BYTES", 10):
            self.assertIn("MB", pf.reference_problem(self.tmp, "refs/x.png"))
        self.assertIn("inside the experiment folder", pf.reference_problem(self.tmp / "refs", "../refs/../../x.png"))
        f.write_bytes(b"<html>")
        self.assertIn("not a PNG", pf.reference_problem(self.tmp, "refs/x.png"))


# --------------------------------------------------------------------------- #
# plans, the workbook and the user's cells
# --------------------------------------------------------------------------- #
class PlanAndWorkbookTests(Base):
    def test_plan_refuses_to_overwrite_an_existing_experiment(self):
        exp = make_experiment(self.tmp, PLAN)
        led = lg.load(exp / "ledger.json")
        led["attempts"].append({"batch": 1, "row_id": "r001", "sample": 1, "sample_id": "r001-1",
                                "request_id": "REQ123", "status": "Completed", "model": NB})
        lg.save(exp / "ledger.json", led)
        with self.assertRaises(mx.PlanError) as cm:
            ex.build_experiment(self.tmp / "src" / "plan.json", exp)
        self.assertIn("already holds an experiment", str(cm.exception))
        self.assertEqual(lg.load(exp / "ledger.json")["attempts"][0]["request_id"], "REQ123")
        r = subprocess.run([sys.executable, str(SCRIPTS / "image.py"), "plan", "--plan",
                            str(self.tmp / "src" / "plan.json"), "--out", str(exp)],
                           capture_output=True, text=True, encoding="utf-8", env=ENV)
        self.assertEqual(r.returncode, 2)
        self.assertNotIn("Traceback", r.stderr)

    def test_clearing_model_reference_and_take_cells_clears_the_ledger_and_stays_cleared(self):
        import openpyxl
        exp = make_experiment(self.tmp, PLAN)
        xl = exp / "experiment.xlsx"

        def set_cells(model, ref, take):
            wb = openpyxl.load_workbook(xl)
            ws = wb["Experiment"]
            head = [c.value for c in ws[1]]
            for col, v in (("Model", model), ("Reference", ref), ("Take", take)):
                ws.cell(row=2, column=head.index(col) + 1).value = v
            wb.save(xl)

        set_cells(NB, "ref1", 3)
        ex.refresh(exp)
        row = lg.load(exp / "ledger.json")["rows"][0]
        self.assertEqual((row["model"], row["reference"], row["take"]), (NB, "ref1", 3))
        set_cells(None, None, None)                                            # the user blanks them
        ex.refresh(exp)
        row = lg.load(exp / "ledger.json")["rows"][0]
        self.assertEqual((row["model"], row["reference"], row["take"]), (None, None, 1))
        ws = openpyxl.load_workbook(xl)["Experiment"]                          # and the new workbook does not write them back
        head = [c.value for c in ws[1]]
        self.assertIn(ws.cell(row=2, column=head.index("Model") + 1).value, (None, ""))
        self.assertIn(ws.cell(row=2, column=head.index("Reference") + 1).value, (None, ""))

    def test_text_that_looks_like_a_formula_or_link_stays_text(self):
        import openpyxl
        exp = make_experiment(self.tmp, PLAN)
        led = lg.load(exp / "ledger.json")
        led["rows"][0]["notes"] = '=HYPERLINK("http://evil.example","x")'
        led["rows"][0]["params"]["camera"] = "=1+1"
        led["rows"][1]["notes"] = "http://evil.example/page"
        wbk.write_workbook(exp / "experiment.xlsx", led, PLAN, prompts={"r001": "=cmd|' /C calc'!A0"})
        ws = openpyxl.load_workbook(exp / "experiment.xlsx")["Experiment"]
        head = [c.value for c in ws[1]]
        for r, col in ((2, "Notes"), (2, "camera"), (2, "Prompt"), (3, "Notes")):
            cell = ws.cell(row=r, column=head.index(col) + 1)
            self.assertEqual(cell.data_type, "s", f"{col} row {r} became {cell.data_type}")
            self.assertIsNone(cell.hyperlink, f"{col} row {r}")
        self.assertEqual(ws.cell(row=2, column=head.index("Notes") + 1).value, '=HYPERLINK("http://evil.example","x")')

    def test_extra_columns_stay_under_their_own_headers(self):
        import openpyxl
        exp = make_experiment(self.tmp, PLAN)
        led = lg.load(exp / "ledger.json")
        led["rows"][0]["extras"] = {"Colour": "red", "Owner": "bob"}
        led["rows"][1]["extras"] = {"Owner": "alice"}                         # the second row has no Colour
        wbk.write_workbook(exp / "experiment.xlsx", led, PLAN)
        ws = openpyxl.load_workbook(exp / "experiment.xlsx")["Experiment"]
        head = [c.value for c in ws[1]]
        owner, colour = head.index("Owner") + 1, head.index("Colour") + 1
        self.assertEqual(ws.cell(row=2, column=colour).value, "red")
        self.assertEqual(ws.cell(row=3, column=owner).value, "alice")
        self.assertIsNone(ws.cell(row=3, column=colour).value)                # not "alice" under Colour

    def test_value_cells_are_text_so_excel_does_not_turn_three_quarter_into_a_date(self):
        import openpyxl
        exp = make_experiment(self.tmp, PLAN)
        ws = openpyxl.load_workbook(exp / "experiment.xlsx")["Experiment"]
        head = [c.value for c in ws[1]]
        for col in ("camera", "lighting", "Model", "Reference", "Notes"):
            self.assertEqual(ws.cell(row=2, column=head.index(col) + 1).number_format, "@", col)
        self.assertNotEqual(ws.cell(row=2, column=head.index("Take") + 1).number_format, "@")

    def test_hostile_numbers_do_not_crash_the_reader_or_the_size_picker(self):
        for v in (float("inf"), float("-inf"), float("nan"), 1e999):
            self.assertIsInstance(wbk._qty(v), (int, float))                    # kept raw, preflight reports it
        self.assertIsNone(pf.aspect_to_wxh("nan:1"))
        self.assertIsNone(pf.aspect_to_wxh("inf:1"))
        self.assertIsNotNone(pf.aspect_to_wxh("4:5"))

    def test_malformed_plans_are_plan_errors_not_tracebacks(self):
        base = copy.deepcopy(PLAN)
        cases = {
            "a dimension called model": lambda p: p["dimensions"].update(model=["slim", "curvy"]),
            "a dimension called Notes": lambda p: p["dimensions"].update(Notes=["a", "b"]),
            "a reference without an id": lambda p: p["references"][0].pop("id"),
            "references holding text": lambda p: p.update(references=["x"]),
            "fixed as a list": lambda p: p.update(fixed=["a"]),
            "prompt as text": lambda p: p.update(prompt="hello"),
            "constraints null": lambda p: p.update(constraints=None),
            "an absolute reference path": lambda p: p["references"][0].update(file="C:/Users/x/other.jpg"),
            "a reference path that climbs out": lambda p: p["references"][0].update(file="../../other.jpg"),
        }
        for name, mutate in cases.items():
            plan = copy.deepcopy(base)
            mutate(plan)
            errs = mx.validate_plan(plan)                                        # must not raise
            self.assertTrue(errs, name)
        for bad in ("no", "", "yes please"):
            self.assertFalse(mx.consent_ok({"consent_acknowledged": bad}), bad)
        self.assertTrue(mx.consent_ok({"consent_acknowledged": "2026-10-07T10:00:00"}))

    def test_an_unknown_intent_is_refused_instead_of_silently_becoming_generic(self):
        plan = copy.deepcopy(PLAN)
        plan["intent"] = "make it pretty"
        with self.assertRaises(mx.PlanError):
            ex.check_intent(plan)
        ex.check_intent(PLAN)

    def test_custom_means_not_in_the_plan_and_plain_plan_values_are_not_custom(self):
        plan = {"brief": "b", "intent": "generic", "dimensions": {"camera": ["front", "side"],
                                                                  "lighting": ["soft daylight"]}}
        cat = cp.load_catalog()
        plain = cp.compile_prompt(plan, {"camera": "front", "lighting": "soft daylight"}, "text", cat)
        self.assertEqual(plain["custom"], [])                                   # "front" is in the plan
        self.assertEqual(plain["no_wording"], [("camera", "front")])            # ...it just has no wording
        typed = cp.compile_prompt(plan, {"camera": "front", "lighting": "overcast daylight"}, "text", cat)
        self.assertEqual(typed["custom"], [("lighting", "overcast daylight")])  # in the catalog, but not in the plan

    def test_a_weak_catalog_value_is_flagged_with_a_reference(self):
        plan = copy.deepcopy(PLAN)
        plan.pop("prompt")
        plan["dimensions"]["camera"] = ["eye-level front", "three-quarter high"]
        exp = make_experiment(self.tmp, plan)
        ledger_select(exp, 6)
        rep = pf.run_preflight(exp, advisor=ADVISOR)
        self.assertTrue(any('"three-quarter high" is weak with a reference image' in i for i in rep["info"]), rep["info"])

    def test_a_newer_refresh_copy_of_the_workbook_is_called_out(self):
        exp = make_experiment(self.tmp, PLAN)
        alt = exp / "experiment.refresh-20260101-000000.xlsx"
        alt.write_bytes((exp / "experiment.xlsx").read_bytes())
        os.utime(alt, (time.time() + 60, time.time() + 60))
        rep = pf.run_preflight(exp, advisor=ADVISOR)
        self.assertTrue(any("NOT read" in w for w in rep["warnings"]))

    def test_the_dry_run_prices_the_model_the_plan_names_not_the_advisors_pick(self):
        plan = copy.deepcopy(PLAN)
        plan["references"], plan.pop("prompt", None)
        plan["references"] = []
        plan["model_strategy"] = "fixed:google/gemini-3.1-flash-image"
        src = self.tmp / "plan.json"
        src.write_text(json.dumps(plan), encoding="utf-8")
        s = ex.build_experiment(src, self.tmp / "o", dry_run=True)
        self.assertEqual(s["model"], "google/gemini-3.1-flash-image")
        self.assertIn("gemini-3.1-flash-image", ex.format_plan_summary(s))
        plan["model_strategy"] = "fixed:no/such-model"
        src.write_text(json.dumps(plan), encoding="utf-8")
        with self.assertRaises(mx.PlanError):
            ex.build_experiment(src, self.tmp / "o", dry_run=True)

    def test_spread_on_a_reference_plan_only_offers_reference_capable_models(self):
        plan = copy.deepcopy(PLAN)
        plan["model_strategy"] = "spread"
        src = self.tmp / "plan.json"
        src.write_text(json.dumps(plan), encoding="utf-8")
        s = ex.build_experiment(src, self.tmp / "o", dry_run=True)
        models = {m for combo in s["result"]["rows"] for m in [dict(zip(s["dims"], combo))["model"]]}
        support = pf.load_reference_support()
        self.assertTrue(models <= {m for m, v in support.items() if v.get("state") in pf.REF_OK})


# --------------------------------------------------------------------------- #
# the gateway layer and downloads
# --------------------------------------------------------------------------- #
class _Handler(http.server.BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_GET(self):
        if self.path == "/good.png":
            self.send_response(200)
            self.send_header("Content-Length", str(len(PNG)))
            self.end_headers()
            self.wfile.write(PNG)
        elif self.path == "/trunc.png":                       # promises more than it sends, then hangs up
            self.send_response(200)
            self.send_header("Content-Length", "5008")
            self.end_headers()
            self.wfile.write(PNG)
        elif self.path == "/oops":
            body = b"oops! an error page"
            self.send_response(200)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        elif self.path == "/redir":
            self.send_response(302)
            self.send_header("Location", "http://127.0.0.1:1/x")
            self.end_headers()
        else:
            self.send_response(404)
            self.end_headers()

    def do_POST(self):
        n = int(self.headers.get("Content-Length") or 0)
        if n:
            self.rfile.read(n)                       # read the body first, or Windows may reset the connection
        self.do_GET()


class DownloadAndGatewayTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _Handler)
        cls.base = f"http://127.0.0.1:{cls.srv.server_address[1]}"
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()

    def setUp(self):
        import tempfile
        self.tmp = Path(tempfile.mkdtemp())
        self._patch = [mock.patch.object(il, "RETRY_SLEEP_SECONDS", 0), mock.patch.object(il, "DOWNLOAD_RETRIES", 1)]
        for p in self._patch:
            p.start()

    def tearDown(self):
        for p in self._patch:
            p.stop()

    def test_a_complete_download_is_renamed_into_place_with_no_temp_file_left(self):
        dest = self.tmp / "a.png"
        il._download(self.base + "/good.png", dest)
        self.assertEqual(dest.read_bytes(), PNG)
        self.assertEqual([p.name for p in self.tmp.iterdir()], ["a.png"])

    def test_a_truncated_download_is_refused_and_leaves_nothing_behind(self):
        dest = self.tmp / "a.png"
        with self.assertRaises(Exception):
            il._download(self.base + "/trunc.png", dest)
        self.assertEqual(list(self.tmp.iterdir()), [])

    def test_an_error_page_never_replaces_an_existing_good_image(self):
        dest = self.tmp / "a.png"
        dest.write_bytes(PNG)
        with self.assertRaises(OSError):
            il._download(self.base + "/oops", dest)
        self.assertEqual(dest.read_bytes(), PNG)
        self.assertEqual([p.name for p in self.tmp.iterdir()], ["a.png"])

    def test_the_gateway_layer_never_follows_a_redirect_with_the_api_key(self):
        with mock.patch.object(il, "GATEWAY", self.base):
            st, body, err = il._req("POST", "/redir", "secret-token", {"x": 1})
        self.assertEqual(st, 302)                                              # surfaced, not followed to another host

    def test_a_bad_token_is_refused_without_being_echoed(self):
        buf = io.StringIO()
        with contextlib.redirect_stderr(buf), self.assertRaises(SystemExit):
            il._clean_token("abc123SECRET\ndef", "$LOOMLOOM_TOKEN_COGFOUNDRY")
        self.assertNotIn("SECRET", buf.getvalue())
        self.assertEqual(il._clean_token("  sk-abc.DEF_123  ", "x"), "sk-abc.DEF_123")

    def test_another_providers_profile_token_is_never_sent_to_this_gateway(self):
        cfgdir = self.tmp / "loomloom"
        cfgdir.mkdir()
        (cfgdir / "config.json").write_text(json.dumps({"active_server": "other", "servers": [
            {"name": "other", "platform": "other", "server": "https://other.example/v1", "token": "THEIR-KEY"}]}),
            encoding="utf-8")
        env = {"APPDATA": str(self.tmp), "HOME": str(self.tmp), "USERPROFILE": str(self.tmp)}
        with mock.patch.dict(os.environ, env), mock.patch.object(Path, "home", lambda: self.tmp):
            for k in ("LOOMLOOM_TOKEN_COGFOUNDRY", "LOOMLOOM_TOKEN"):
                os.environ.pop(k, None)
            buf = io.StringIO()
            with contextlib.redirect_stderr(buf), self.assertRaises(SystemExit):
                il.token()
            self.assertNotIn("THEIR-KEY", buf.getvalue())
            (cfgdir / "config.json").write_text(json.dumps({"active_server": "cogfoundry", "servers": [
                {"name": "cogfoundry", "platform": "cogfoundry", "server": "https://loomloom.cogfoundry.ai/v1",
                 "token": "OUR-KEY"}]}), encoding="utf-8")
            self.assertEqual(il.token(), "OUR-KEY")

    def test_run_alloc_cannot_be_combined_with_the_v2_flags(self):
        r = subprocess.run([sys.executable, str(SCRIPTS / "image.py"), "run", "--alloc", "m:1:default:-",
                            "--prompt", "p", "--confirm", "x", "--dir", str(self.tmp)],
                           capture_output=True, text=True, encoding="utf-8", env=ENV)
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("cannot be combined", r.stderr)


# --------------------------------------------------------------------------- #
# the contact sheet and the page builder
# --------------------------------------------------------------------------- #
class SheetSafetyTests(RunBase):
    def done(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway())

    def test_only_relative_image_paths_inside_the_experiment_are_ever_shown_or_embedded(self):
        self.done()
        secret = self.tmp / "secret.png"
        secret.write_bytes(PNG)
        led = self.ledger()
        for bad in ("../secret.png", str(secret), "https://attacker.example/b.png", "round-1/notes.txt",
                    "file:///etc/passwd"):
            led["attempts"][0]["file"] = bad
            lg.save(self.exp / "ledger.json", led)
            for inline in (False, True):
                d = sh.build_data(self.exp, inline=inline)
                smp = d["rows"][0]["samples"][0]
                self.assertIsNone(smp["file"], bad)
                self.assertIn("unsafe", smp["error"], bad)
        led["attempts"][0]["file"] = "round-1/gone.png"                           # a missing image is reported
        lg.save(self.exp / "ledger.json", led)
        self.assertIn("missing", sh.build_data(self.exp, inline=True)["rows"][0]["samples"][0]["error"])

    def test_links_are_relative_to_where_the_page_is_written(self):
        self.done()
        out = self.tmp / "elsewhere" / "deep" / "sheet.html"
        out.parent.mkdir(parents=True)
        d = sh.build_data(self.exp, out_dir=out.parent)
        rel = d["rows"][0]["samples"][0]["file"]
        self.assertTrue((out.parent / rel).resolve().exists(), rel)

    def test_an_unreadable_plan_or_a_flagged_snapshot_counts_as_a_person_experiment(self):
        self.done()
        self.assertFalse(sh.build_data(self.exp)["person"])
        plan = self.exp / "plan.json"
        good = plan.read_text(encoding="utf-8")
        plan.write_text("{not json", encoding="utf-8")                              # corrupt: fail closed
        self.assertTrue(sh.build_data(self.exp)["person"])
        with self.assertRaises(sh.SheetError):
            sh.build_data(self.exp, inline=True)
        plan.unlink()                                                                # missing, but the snapshot says person
        snap = next((self.exp / "snapshots").glob("*.json"))
        data = json.loads(snap.read_text(encoding="utf-8"))
        data["rows"][0]["references"][0]["person"] = True
        snap.write_text(json.dumps(data), encoding="utf-8")
        self.assertTrue(sh.build_data(self.exp)["person"])
        plan.write_text(good, encoding="utf-8")

    def test_the_page_builder_refuses_a_folder_whose_plan_cannot_be_read(self):
        self.done()
        (self.exp / "plan.json").write_text("{not json", encoding="utf-8")
        r = subprocess.run([sys.executable, str(SCRIPTS / "build-exploration-page.py"), "--session", str(self.exp),
                            "--title", "t", "--subject", "s", "--invocation", "i"],
                           capture_output=True, text=True, encoding="utf-8", env=ENV)
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("cannot be read", r.stderr)


# --------------------------------------------------------------------------- #
# model catalog corrections found by reading the live model pages (2026-10-07)
# --------------------------------------------------------------------------- #
class ModelCatalogTests(Base):
    SUN = "openai/gpt-image-2.5-sunburst"

    def test_the_live_pages_corrected_sunburst_flare_prices_and_prompt_limits(self):
        cat = pf.il.load_model_catalog()
        for m in (self.SUN, "openai/gpt-image-2.5-flare"):
            req = cat[m]["request"]
            self.assertEqual((req["size_param"], req["size_step"]), ("wxh", "16"), m)   # they DO take a size
        self.assertGreater(cat["google/gemini-3-pro-image"]["usd_per_image"], 0.1)       # was off by 10x
        self.assertLessEqual(cat["bytedance/doubao-seedream-5-0-pro"]["usd_per_image"], 0.06)
        for m in ("bytedance/doubao-seedream-5-0-pro", "bytedance/doubao-seedream-5-0-lite",
                  "bytedance/doubao-seedream-4.5"):
            self.assertEqual(cat[m]["prompt_max_chars"], 500, m)

    def test_an_exact_aspect_ratio_is_requested_from_models_that_take_any_size(self):
        for ratio in ("4:5", "1:1", "16:9", "3:2", "9:16", "5:4"):
            tok = pf.exact_wxh(ratio, 1024 * 1024)
            w, h = (int(x) for x in tok.split("x"))
            a, b = (int(x) for x in ratio.split(":"))
            self.assertEqual((w % 16, h % 16), (0, 0), tok)
            self.assertEqual(w * b, h * a, tok)                       # exactly the ratio, not the nearest
            self.assertGreaterEqual(w * h, 1024 * 1024, tok)
            self.assertLessEqual(max(w, h), 3840)
        self.assertIsNone(pf.exact_wxh("4:1"))                       # beyond the documented 3:1
        self.assertIsNone(pf.exact_wxh("nan:1"))

    def test_sunburst_gets_the_exact_4_5_size_and_models_without_a_size_step_keep_the_old_behavior(self):
        adv = pf.Advisor("poster / flyer")
        s = adv.size_for(self.SUN, "4:5")
        w, h = (int(x) for x in s["token"].split("x"))
        self.assertEqual(w * 5, h * 4)
        self.assertGreaterEqual(w * h, 1536 * 1024)                  # about the poster intent's own size, not a thumbnail
        self.assertIsNone(s["aspect_ratio"])
        g2 = adv.size_for("openai/gpt-image-2", "4:5")               # no documented step: unchanged, so not exact
        self.assertNotEqual(int(g2["token"].split("x")[0]) * 5, int(g2["token"].split("x")[1]) * 4)
        nb = adv.size_for("google/gemini-3.1-flash-image", "4:5")    # a real ratio field
        self.assertEqual(nb["aspect_ratio"], "4:5")

    def test_the_size_reaches_the_gateway_request(self):
        adv = pf.Advisor("poster / flyer")
        tok = adv.size_for(self.SUN, "4:5")["token"]
        body = gen.build_request({"id": "r001", "model": self.SUN, "prompt": "p", "size": tok, "aspect_ratio": None,
                                  "references": []}, pf.il.load_model_catalog(), self.tmp)
        self.assertEqual(body["size"], tok)

    def test_a_prompt_longer_than_a_models_limit_is_refused_or_routed_around(self):
        plan = copy.deepcopy(PLAN)
        plan["references"] = []
        plan.pop("prompt", None)
        plan["prompt"] = {"text_prefix": "x" * 800}
        plan["model_strategy"] = "fixed:bytedance/doubao-seedream-5-0-lite"
        src = self.tmp / "plan.json"
        src.write_text(json.dumps(plan), encoding="utf-8")
        s = ex.build_experiment(src, self.tmp / "o", dry_run=True)
        self.assertTrue(any("accepts at most 500" in w for w in s["warnings"]), s["warnings"])
        # at preflight an explicit Seedream choice is an issue for every row ...
        exp = make_experiment(self.tmp / "e", plan)
        ledger_select(exp, 2)
        rep = pf.run_preflight(exp, advisor=ADVISOR)
        self.assertEqual(rep["ready"], 0)
        self.assertTrue(all("accepts at most 500" in p for i in rep["issues"] for p in i["problems"]))
        # ... and with no explicit model the Advisor never auto-picks a model that would refuse the prompt
        plan["model_strategy"] = "single"
        exp2 = make_experiment(self.tmp / "e2", plan)
        ledger_select(exp2, 2)
        rep2 = pf.run_preflight(exp2, advisor=ADVISOR)
        self.assertEqual(rep2["ready"], 2)
        self.assertFalse(any(m.startswith("bytedance/") for m in rep2["models"]))


# --------------------------------------------------------------------------- #
# the quality setting (adidas run: one model billed $0.030 or $0.055 per image on quality=auto)
# --------------------------------------------------------------------------- #
class QualityTests(Base):
    SUN = "openai/gpt-image-2.5-sunburst"

    def text_plan(self, quality=None, strategy=None):
        plan = copy.deepcopy(PLAN)
        plan["references"] = []
        plan.pop("prompt", None)
        plan["model_strategy"] = strategy or ("fixed:" + self.SUN)
        if quality:
            plan["quality"] = quality
        return plan

    def test_catalog_lists_the_quality_values_of_the_models_that_have_them(self):
        cat = pf.il.load_model_catalog()
        self.assertEqual(pf.il._shape(cat[self.SUN])["quality_values"], ["auto", "low", "medium", "high", "xhigh", "max"])
        self.assertIn("high", pf.il._shape(cat["openai/gpt-image-2"])["quality_values"])
        self.assertEqual(pf.il._shape(cat["google/gemini-3.1-flash-image"])["quality_values"], [])

    def test_quality_must_be_text(self):
        plan = self.text_plan()
        plan["quality"] = 5
        self.assertTrue(any("quality" in e for e in mx.validate_plan(plan)))
        plan["quality"] = "  "
        self.assertTrue(any("quality" in e for e in mx.validate_plan(plan)))
        plan["quality"] = "medium"
        self.assertEqual(mx.validate_plan(plan), [])

    def test_the_quality_reaches_the_snapshot_and_the_request_and_the_price_is_honestly_unknown(self):
        exp = make_experiment(self.tmp, self.text_plan("medium"))
        ledger_select(exp, 3)
        rep = pf.run_preflight(exp, advisor=ADVISOR)
        self.assertEqual(rep["ready"], 3)
        rows = rep["snapshot"]["rows"]
        self.assertTrue(all(r["quality"] == "medium" for r in rows))
        self.assertEqual(len(rep["unverified_rows"]), 3)          # the catalog's flat price is NOT used for an explicit quality
        self.assertEqual(rep["known_usd"], 0.0)
        body = gen.build_request(rows[0], pf.il.load_model_catalog(), exp)
        self.assertEqual(body["quality"], "medium")
        with self.assertRaises(gen.RunRefused):                     # no estimate, so the user must set a limit
            gen.run_batch(exp, rep["fingerprint"], gateway=FakeGateway(), log=quiet, record_prices=False)

    def test_auto_or_no_quality_sends_nothing_and_keeps_the_old_estimate(self):
        for q in (None, "auto"):
            self.tearDown()
            self.setUp()
            exp = make_experiment(self.tmp, self.text_plan(q))
            ledger_select(exp, 2)
            rep = pf.run_preflight(exp, advisor=ADVISOR)
            self.assertTrue(all("quality" not in r for r in rep["snapshot"]["rows"]), q)
            self.assertEqual(len(rep["unverified_rows"]), 0, q)
            body = gen.build_request(rep["snapshot"]["rows"][0], pf.il.load_model_catalog(), exp)
            self.assertNotIn("quality", body)

    def test_a_model_without_that_quality_is_an_issue_and_an_auto_pick_avoids_it(self):
        exp = make_experiment(self.tmp, self.text_plan("medium", "fixed:google/gemini-3.1-flash-image"))
        ledger_select(exp, 2)
        rep = pf.run_preflight(exp, advisor=ADVISOR)
        self.assertEqual(rep["ready"], 0)
        self.assertTrue(all("no quality setting" in p for i in rep["issues"] for p in i["problems"]))
        exp2 = make_experiment(self.tmp / "b", self.text_plan("high", "single"))
        ledger_select(exp2, 2)
        rep2 = pf.run_preflight(exp2, advisor=ADVISOR)
        self.assertEqual(rep2["ready"], 2)
        cat = pf.il.load_model_catalog()
        self.assertTrue(all("high" in pf.il._shape(cat[m])["quality_values"] for m in rep2["models"]))

    def test_the_price_observed_at_a_quality_is_used_for_the_next_estimate_of_that_quality_only(self):
        exp = make_experiment(self.tmp, self.text_plan("medium"))
        ledger_select(exp, 2)
        rep = pf.run_preflight(exp, advisor=ADVISOR)
        obs_file = self.tmp / "price-observations.json"
        old, old_support = gen.il.PRICE_OBSERVATIONS_FILE, pf.REF_SUPPORT_FILE
        gen.il.PRICE_OBSERVATIONS_FILE = str(obs_file)
        try:
            gen.run_batch(exp, rep["fingerprint"], gateway=FakeGateway(cost=0.0123), poll_seconds=0.01, max_usd=1, log=quiet)
        finally:
            gen.il.PRICE_OBSERVATIONS_FILE = old
        obs = json.loads(obs_file.read_text(encoding="utf-8"))
        key = next(k for k in obs if "|q=medium" in k)
        self.assertAlmostEqual(obs[key], 0.0123, places=4)
        adv = pf.Advisor("product / e-commerce shot")
        adv.observations = dict(obs)
        size = {"token": key.split("|")[1] if key.split("|")[1] != "default" else None, "px": 1024 * 1024}
        self.assertEqual(adv.price(self.SUN, size, "text", {}, "medium")[1], "observed")
        self.assertEqual(adv.price(self.SUN, size, "text", {}, "high")[1], "unverified")   # another quality: still unknown

    def test_the_dry_run_shows_the_quality_and_why_the_price_is_unknown(self):
        src = self.tmp / "plan.json"
        src.write_text(json.dumps(self.text_plan("high")), encoding="utf-8")
        text = ex.format_plan_summary(ex.build_experiment(src, self.tmp / "o", dry_run=True))
        self.assertIn("Quality:             high", text)
        self.assertIn("no observed price for this model, size and quality yet", text)
        bad = self.text_plan("high", "fixed:google/gemini-3.1-flash-image")
        src.write_text(json.dumps(bad), encoding="utf-8")
        with self.assertRaises(mx.PlanError):
            ex.build_experiment(src, self.tmp / "o", dry_run=True)


if __name__ == "__main__":
    unittest.main()
