import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "scripts"))
import ledger as lg  # noqa: E402
import matrix as mx  # noqa: E402
import workbook as wbk  # noqa: E402

PLAN = json.loads((HERE / "fixtures" / "plan-mighty.json").read_text(encoding="utf-8"))
NAMES = ["camera", "lighting"]


def make_ledger(n=4):
    dims = mx.dimensions(PLAN)
    rows = mx.cover(dims, [], target=n)["rows"][:n]
    return lg.new_ledger(PLAN, NAMES, rows)


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.path = self.tmp / "experiment.xlsx"
        self.ledger = make_ledger()
        wbk.write_workbook(self.path, self.ledger, PLAN)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def edit(self, fn):
        """Simulate the user: open the workbook, change cells, save (test only)."""
        import openpyxl
        wb = openpyxl.load_workbook(self.path)
        fn(wb["Experiment"])
        wb.save(self.path)


class WriteRead(Base):
    def test_roundtrip(self):
        got = wbk.read_workbook(self.path, NAMES)
        self.assertEqual(len(got["rows"]), 4)
        r0 = got["rows"][0]
        self.assertEqual(r0["id"], "r001")
        self.assertIs(r0["selected"], True)                    # native checkbox reads back as a bool
        self.assertEqual(r0["params"], self.ledger["rows"][0]["params"])
        self.assertEqual(r0["take"], 1)
        self.assertEqual(got["extra_columns"], [])

    def test_an_older_workbook_with_qty_and_image_columns_is_read_without_extras(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.path)
        ws = wb["Experiment"]
        head = [c.value for c in ws[1]]
        ws.cell(row=1, column=head.index("Take") + 1).value = "Qty"
        ws.cell(row=1, column=head.index("File") + 1).value = "Images"
        ws.cell(row=2, column=head.index("Qty" if "Qty" in head else "Take") + 1).value = 3
        ws.cell(row=1, column=ws.max_column + 1).value = "Image 1"
        wb.save(self.path)
        got = wbk.read_workbook(self.path, NAMES)
        self.assertEqual(got["extra_columns"], [])
        self.assertNotIn("take", got["rows"][0])
        self.assertEqual(got["rows"][0]["legacy_qty"], 3)
        warnings = lg.merge_workbook(self.ledger, got["rows"])
        self.assertTrue(any("Qty 3 is no longer used" in w and "add-takes" in w for w in warnings))
        self.assertEqual(self.ledger["rows"][0]["take"], 1)            # the ledger's Take is untouched

    def test_system_columns_are_not_read_back(self):
        got = wbk.read_workbook(self.path, NAMES)
        self.assertNotIn("prompt", got["rows"][0])
        self.assertNotIn("status", got["rows"][0])

    def test_read_me_sheet_has_the_plan_summary(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.path)
        text = "\n".join(str(c.value) for c in wb["Read me"]["A"] if c.value)
        self.assertIn("Variable - camera", text)
        self.assertIn("ID, Prompt, Status", text)

    def test_locked_file_falls_back_to_a_refresh_file(self):
        real = wbk._write
        calls = []

        def flaky(*a, **k):
            calls.append(a[1])
            if len(calls) == 1:
                raise PermissionError("file is open in Excel")
            return real(*a, **k)

        with mock.patch.object(wbk, "_write", flaky):
            out = wbk.write_workbook(self.path, self.ledger, PLAN)
        self.assertIn("refresh-", out.name)
        self.assertTrue(out.exists())

    def test_unwritable_target_falls_back_to_a_refresh_file_for_real(self):
        """No mocks: a target that cannot be opened for writing (a directory here; an Excel
        lock in real life) must raise XlsxWriter's FileCreateError, which we must handle."""
        blocked = self.tmp / "blocked.xlsx"
        blocked.mkdir()
        out = wbk.write_workbook(blocked, self.ledger, PLAN)
        self.assertIn("refresh-", out.name)
        self.assertTrue(out.exists() and out.is_file())

    def test_backup_previous(self):
        prev = wbk.backup_previous(self.path)
        self.assertEqual(prev.name, "experiment.prev.xlsx")
        self.assertTrue(prev.exists())


class Merge(Base):
    def merged(self):
        sheet = wbk.read_workbook(self.path, NAMES)
        warnings = lg.merge_workbook(self.ledger, sheet["rows"])
        return {r["id"]: r for r in self.ledger["rows"]}, warnings

    def test_user_edits_are_merged(self):
        def fn(ws):
            ws["B3"] = False                     # untick r002
            ws["C2"] = "low angle"               # change r001's camera
            ws["F2"] = 3                         # Take of r001
            ws["L2"] = "keep"                    # Notes of r001
        self.edit(fn)
        rows, warnings = self.merged()
        self.assertEqual(warnings, [])
        self.assertFalse(rows["r002"]["selected"])
        self.assertEqual(rows["r001"]["params"]["camera"], "low angle")
        self.assertEqual(rows["r001"]["take"], 3)
        self.assertEqual(rows["r001"]["notes"], "keep")

    def test_edited_id_is_not_recognized_and_the_original_becomes_removed(self):
        self.edit(lambda ws: ws.__setitem__("A2", "my-product-1"))
        rows, warnings = self.merged()
        self.assertTrue(any("not recognized (IDs are read-only)" in w for w in warnings))
        self.assertEqual(rows["r001"]["status"], "Removed")
        self.assertIn("r005", rows)                                  # new row, new system ID

    def test_duplicate_id_gets_a_new_id(self):
        self.edit(lambda ws: ws.__setitem__("A3", "r001"))
        rows, warnings = self.merged()
        self.assertTrue(any("duplicate ID r001" in w for w in warnings))
        self.assertEqual(rows["r002"]["status"], "Removed")          # r002 no longer in the workbook
        self.assertIn("r005", rows)

    def test_row_without_id_is_added_and_deleted_row_is_removed(self):
        def fn(ws):
            ws.delete_rows(5)                                         # r004 gone
            ws.append([None, True, "front", "soft daylight"])         # a new row without an ID
        self.edit(fn)
        rows, warnings = self.merged()
        self.assertEqual(rows["r004"]["status"], "Removed")
        self.assertTrue(any("no ID, assigned r005" in w for w in warnings))
        self.assertEqual(rows["r005"]["params"], {"camera": "front", "lighting": "soft daylight"})

    def test_extra_columns_are_carried_and_ignored(self):
        self.edit(lambda ws: (ws.cell(1, 20, "Wardrobe"), ws.cell(2, 20, "red coat")))
        sheet = wbk.read_workbook(self.path, NAMES)
        self.assertEqual(sheet["extra_columns"], ["Wardrobe"])
        lg.merge_workbook(self.ledger, sheet["rows"])
        self.assertEqual(self.ledger["rows"][0]["extras"], {"Wardrobe": "red coat"})
        # a new dimension column is NOT a plan dimension
        self.assertEqual(self.ledger["experiment"]["dimensions"], NAMES)

    def test_ids_are_never_reused(self):
        self.edit(lambda ws: ws.delete_rows(5))
        self.merged()
        again = lg.add_row(self.ledger, {"camera": "x", "lighting": "y"})
        self.assertEqual(again["id"], "r005")


class Refresh(Base):
    def test_regenerate_keeps_user_cells_and_a_prev_copy(self):
        self.edit(lambda ws: ws.__setitem__("F3", 2))                # user sets Take 2 on r002
        # the documented refresh: read -> merge -> backup -> write fresh
        sheet = wbk.read_workbook(self.path, NAMES)
        lg.merge_workbook(self.ledger, sheet["rows"])
        wbk.backup_previous(self.path)
        wbk.write_workbook(self.path, self.ledger, PLAN, prompts={"r001": "hello prompt"})
        again = wbk.read_workbook(self.path, NAMES)
        self.assertEqual(again["rows"][1]["take"], 2)                # user edit survived the regenerate
        self.assertTrue((self.tmp / "experiment.prev.xlsx").exists())
        import openpyxl
        ws = openpyxl.load_workbook(self.path)["Experiment"]
        self.assertEqual(ws["H2"].value, "hello prompt")             # system Prompt column written


class Thumbnails(unittest.TestCase):
    """Auto filter, in-cell thumbnails with a click link, and the hover preview."""

    def setUp(self):
        from PIL import Image
        self.tmp = Path(tempfile.mkdtemp())
        self.path = self.tmp / "experiment.xlsx"
        (self.tmp / "round-1").mkdir()
        for n in ("a", "b", "c", "d"):
            Image.new("RGB", (400, 500), (200, 40, 40)).save(self.tmp / "round-1" / f"{n}.png")
        self.ledger = make_ledger()
        # r001 has an older image too (it was regenerated after an edit); the row shows its latest, d.png
        self.links = {"r001": ["external:round-1/a.png", "external:round-1/d.png"],
                      "r002": ["external:round-1/b.png"], "r003": ["external:round-1/c.png"]}

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def parts(self):
        import zipfile
        z = zipfile.ZipFile(self.path)
        return z, z.namelist()

    def test_auto_filter_on_the_header_row(self):
        wbk.write_workbook(self.path, self.ledger, PLAN)
        import openpyxl
        ws = openpyxl.load_workbook(self.path)["Experiment"]
        self.assertTrue(ws.auto_filter.ref.startswith("A1:"))
        self.assertEqual(ws.freeze_panes, "C2")

    def test_thumbnails_hover_and_click_link(self):
        import re
        wbk.write_workbook(self.path, self.ledger, PLAN, image_links=self.links)
        z, names = self.parts()
        self.assertIsNone(z.testzip())
        self.assertTrue(any("richData" in n for n in names))                    # in-cell pictures
        vml = z.read("xl/drawings/vmlDrawing1.vml").decode()
        self.assertEqual(len(re.findall(r'<v:fill o:relid="rId\d+"', vml)), 3)    # one picture per row
        self.assertIn('Extension="jpeg"', z.read("[Content_Types].xml").decode())
        rels = z.read("xl/drawings/_rels/vmlDrawing1.vml.rels").decode()
        self.assertEqual(rels.count("<Relationship "), 3)
        for n in re.findall(r"Target=\"\.\./(media/[^\"]+)\"", rels):
            self.assertIn("xl/" + n, names)
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        self.assertIn("round-1/d.png", z.read("xl/worksheets/_rels/sheet1.xml.rels").decode())   # click opens the raw file
        rels_text = z.read("xl/worksheets/_rels/sheet1.xml.rels").decode()
        self.assertNotIn("round-1/a.png", rels_text)                      # r001's older image is not what its picture links to

    def test_one_image_per_row_shown_by_name_with_a_link(self):
        wbk.write_workbook(self.path, self.ledger, PLAN, image_links=self.links)
        got = wbk.read_workbook(self.path, NAMES)
        self.assertEqual(got["extra_columns"], [])                  # Image and File are system columns, not user extras
        import openpyxl
        ws = openpyxl.load_workbook(self.path)["Experiment"]
        head = [c.value for c in ws[1]]
        self.assertEqual([h for h in head if h and (h.startswith("Image") or h == "File")], ["Image", "File"])
        cell = ws.cell(row=2, column=head.index("File") + 1)
        self.assertEqual(cell.value, "d.png")                       # the latest image, by name
        self.assertIn("d.png", cell.hyperlink.target or cell.hyperlink.location or "")
        self.assertIsNone(ws.cell(row=5, column=head.index("File") + 1).value)     # r004 has no image yet

    def test_thumbnail_and_hover_are_small(self):
        wbk.write_workbook(self.path, self.ledger, PLAN, image_links=self.links)
        self.assertLess(self.path.stat().st_size, 400_000)

    def test_regenerating_keeps_user_edits(self):
        import openpyxl
        wbk.write_workbook(self.path, self.ledger, PLAN, image_links=self.links)
        wb = openpyxl.load_workbook(self.path)
        ws = wb["Experiment"]
        head = [c.value for c in ws[1]]
        ws.cell(row=2, column=head.index("Notes") + 1).value = "keep me"
        ws.cell(row=2, column=head.index("camera") + 1).value = "low angle"
        wb.save(self.path)                                          # the user edits and saves (this drops the pictures, as Excel would not)
        lg.merge_workbook(self.ledger, wbk.read_workbook(self.path, NAMES)["rows"])
        wbk.write_workbook(self.path, self.ledger, PLAN, image_links=self.links)       # the regenerate
        again = wbk.read_workbook(self.path, NAMES)["rows"][0]
        self.assertEqual((again["notes"], again["params"]["camera"]), ("keep me", "low angle"))

    def test_path_outside_the_experiment_folder_is_not_embedded(self):
        outside = self.tmp.parent / "outside-secret.png"
        from PIL import Image
        Image.new("RGB", (50, 50)).save(outside)
        try:
            wbk.write_workbook(self.path, self.ledger, PLAN, image_links={"r001": ["external:../outside-secret.png"]})
        finally:
            outside.unlink()
        z, names = self.parts()
        self.assertFalse(any("media" in n for n in names))

    def test_broken_or_missing_image_is_skipped_not_fatal(self):
        (self.tmp / "round-1" / "bad.png").write_bytes(b"not a png")
        wbk.write_workbook(self.path, self.ledger, PLAN, image_links={
            "r001": ["external:round-1/bad.png", "external:round-1/missing.png", "external:round-1/a.png"]})
        z, names = self.parts()
        self.assertIsNone(z.testzip())
        self.assertEqual(len([n for n in names if n.startswith("xl/media/hover")]), 1)

    def test_hover_box_has_the_pictures_shape_and_no_cell_anchor(self):
        """Regression: the note picture looked stretched. Excel sized the note from the cells it spans (tall wrapped
        rows made it taller than planned) and a fixed 4:5 box stretched square or wide images."""
        import re
        from PIL import Image
        Image.new("RGB", (1000, 500), (10, 10, 200)).save(self.tmp / "round-1" / "wide.png")
        Image.new("RGB", (500, 500), (10, 200, 10)).save(self.tmp / "round-1" / "square.png")
        links = {"r001": ["external:round-1/a.png"], "r002": ["external:round-1/wide.png"],
                 "r003": ["external:round-1/square.png"]}
        wbk.write_workbook(self.path, self.ledger, PLAN, image_links=links)
        z, _ = self.parts()
        vml = z.read("xl/drawings/vmlDrawing1.vml").decode()
        self.assertNotIn("<x:Anchor>", vml)
        sizes = [(float(w), float(h)) for w, h in re.findall(r"width:([\d.]+)pt;height:([\d.]+)pt", vml)]
        self.assertEqual(len(sizes), 3)
        ratios = sorted(round(w / h, 2) for w, h in sizes)
        self.assertEqual(ratios, [0.8, 1.0, 2.0])                    # 4:5, 1:1 and 2:1 as the pictures are

    def test_failed_hover_injection_leaves_a_valid_workbook(self):
        real = os.replace

        def replace(src, dst):                         # fail only the hover step's own temp file
            if str(src).endswith(".tmp.tmp"):
                raise OSError("boom")
            return real(src, dst)

        with mock.patch.object(wbk.os, "replace", side_effect=replace):
            wbk.write_workbook(self.path, self.ledger, PLAN, image_links=self.links)
        self.assertEqual([x.name for x in self.tmp.glob("*.tmp*")], [])
        z, names = self.parts()
        self.assertIsNone(z.testzip())
        self.assertEqual(wbk.read_workbook(self.path, NAMES)["rows"][0]["id"], "r001")

    def test_without_pillow_there_are_no_picture_columns(self):
        with mock.patch.object(wbk, "_have_pillow", return_value=False):
            wbk.write_workbook(self.path, self.ledger, PLAN, image_links=self.links)
        import openpyxl
        head = [c.value for c in openpyxl.load_workbook(self.path)["Experiment"][1]]
        self.assertNotIn("Image", head)
        self.assertIn("File", head)


class ReviewFixTests(Base):
    """Regressions for the review of workbook.py (2026-10-08)."""

    def setUp(self):
        super().setUp()
        from PIL import Image
        (self.tmp / "round-1").mkdir()
        for n in ("a", "b"):
            Image.new("RGB", (400, 500), (200, 40, 40)).save(self.tmp / "round-1" / f"{n}.png")
        self.links = {"r001": ["external:round-1/a.png"]}

    def parts(self):
        import zipfile
        z = zipfile.ZipFile(self.path)
        return z, z.namelist()

    def test_a_truncated_or_empty_cache_entry_is_regenerated_not_fatal(self):
        wbk.write_workbook(self.path, self.ledger, PLAN, image_links=self.links)
        cache = self.tmp / ".thumbs"
        files = sorted(cache.glob("*.jpg"))
        self.assertEqual(len(files), 2)
        files[0].write_bytes(b"")                                   # an interrupted write
        files[1].write_bytes(files[1].read_bytes()[:30])
        wbk.write_workbook(self.path, self.ledger, PLAN, image_links=self.links)       # must not raise
        z, names = self.parts()
        self.assertIsNone(z.testzip())
        self.assertTrue(any("hover" in n for n in names))
        self.assertTrue(all(f.stat().st_size > 100 for f in cache.glob("*.jpg")))     # healed

    def test_the_newest_readable_image_is_used_when_the_latest_is_bad(self):
        (self.tmp / "round-1" / "bad.png").write_bytes(b"not a png")
        wbk.write_workbook(self.path, self.ledger, PLAN,
                           image_links={"r001": ["external:round-1/a.png", "external:round-1/bad.png"]})
        z, names = self.parts()
        self.assertEqual(len([n for n in names if n.startswith("xl/media/hover")]), 1)   # a.png stood in for the bad latest file

    def test_the_file_column_links_the_latest_image_that_exists(self):
        import openpyxl
        wbk.write_workbook(self.path, self.ledger, PLAN,
                           image_links={"r001": ["external:round-1/a.png", "external:round-1/missing.png"]})
        ws = openpyxl.load_workbook(self.path)["Experiment"]
        head = [c.value for c in ws[1]]
        self.assertEqual(ws.cell(row=2, column=head.index("File") + 1).value, "a.png")

    def test_a_failure_while_writing_leaves_the_previous_workbook_untouched(self):
        wbk.write_workbook(self.path, self.ledger, PLAN)
        before = self.path.read_bytes()
        with mock.patch.object(wbk, "_inject_hover", side_effect=ValueError("disk full")):
            with self.assertRaises(ValueError):
                wbk.write_workbook(self.path, self.ledger, PLAN, image_links=self.links)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual([x.name for x in self.tmp.glob("*.tmp*")], [])

    def test_a_date_in_an_extra_column_is_stored_as_text_and_the_ledger_saves(self):
        import datetime
        import openpyxl
        wb = openpyxl.load_workbook(self.path)
        ws = wb["Experiment"]
        col = ws.max_column + 1
        ws.cell(row=1, column=col).value = "Deadline"
        ws.cell(row=2, column=col).value = datetime.datetime(2026, 10, 1, 9, 30)
        wb.save(self.path)
        sheet = wbk.read_workbook(self.path, NAMES)
        self.assertEqual(sheet["rows"][0]["extras"]["Deadline"], "2026-10-01T09:30:00")
        lg.merge_workbook(self.ledger, sheet["rows"])
        lg.save(self.tmp / "ledger.json", self.ledger)                # used to raise TypeError

    def test_renaming_or_deleting_a_dimension_column_does_not_wipe_the_values(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.path)
        ws = wb["Experiment"]
        head = [c.value for c in ws[1]]
        ws.cell(row=1, column=head.index("camera") + 1).value = "Camera"             # capitalisation changed
        wb.save(self.path)
        sheet = wbk.read_workbook(self.path, NAMES)
        self.assertEqual(sheet["rows"][0]["params"]["camera"], self.ledger["rows"][0]["params"]["camera"])
        self.assertEqual(sheet["extra_columns"], [])
        ws.cell(row=1, column=head.index("camera") + 1).value = "something else"      # a different name: the column is gone
        wb.save(self.path)
        sheet = wbk.read_workbook(self.path, NAMES)
        self.assertEqual(sheet["missing_columns"], ["camera"])
        before = dict(self.ledger["rows"][0]["params"])
        lg.merge_workbook(self.ledger, sheet["rows"])
        self.assertEqual(self.ledger["rows"][0]["params"]["camera"], before["camera"])   # kept, not erased

    def test_control_characters_do_not_come_back_as_escape_text(self):
        import openpyxl
        self.ledger["rows"][0]["notes"] = "a\x00b\x0bc"
        wbk.write_workbook(self.path, self.ledger, PLAN)
        ws = openpyxl.load_workbook(self.path)["Experiment"]
        head = [c.value for c in ws[1]]
        self.assertEqual(ws.cell(row=2, column=head.index("Notes") + 1).value, "abc")

    def test_transparent_pixels_become_white_not_their_hidden_colour(self):
        from PIL import Image
        im = Image.new("RGBA", (4, 4), (255, 0, 0, 0))                # fully transparent red
        flat = wbk._flatten(im)
        self.assertEqual(flat.getpixel((1, 1)), (255, 255, 255))


class LedgerIO(unittest.TestCase):
    def test_atomic_save_and_load(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            led = make_ledger()
            lg.save(tmp / "ledger.json", led)
            self.assertFalse((tmp / "ledger.json.tmp").exists())
            self.assertEqual(lg.load(tmp / "ledger.json")["rows"][0]["id"], "r001")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
