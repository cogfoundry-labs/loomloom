"""A row is exactly one image; more images of the same values are more rows with the next Take."""
import copy
import json
import os
import time
import subprocess
import sys
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(HERE))
import experiment as ex  # noqa: E402
import ledger as lg  # noqa: E402
import matrix as mx  # noqa: E402
import preflight as pf  # noqa: E402
import test_generate as tg  # noqa: E402
from test_generate import FakeGateway  # noqa: E402
from test_preflight import ADVISOR, Base, PLAN  # noqa: E402


class TakesTests(Base):
    prepare = tg.GenerateTests.prepare
    run_gen = tg.GenerateTests.run_gen

    def test_add_takes_copies_the_values_and_numbers_the_next_takes(self):
        self.prepare(2)
        res = ex.add_takes(self.exp, ["r001"], 2)
        self.assertEqual([t for _, t in res["added"]], [2, 3])
        led = self.ledger()
        new = [r for r in led["rows"] if r["id"] in {i for i, _ in res["added"]}]
        self.assertTrue(all(r["selected"] and r["status"] == "Draft" for r in new))
        self.assertTrue(all(r["params"] == led["rows"][0]["params"] for r in new))
        again = ex.add_takes(self.exp, ["r001"], 1)                      # numbering continues from the highest take
        self.assertEqual([t for _, t in again["added"]], [4])

    def test_takes_of_a_row_that_already_has_its_image_cost_one_image_each(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway())
        ex.add_takes(self.exp, ["r001"], 2)
        rep = pf.run_preflight(self.exp, advisor=ADVISOR)
        self.assertEqual(rep["images"], 2)                              # r001 itself is skipped; its two takes are new
        self.assertTrue(any("already has its image" in w for w in rep["warnings"]))
        self.run_gen(rep["fingerprint"], FakeGateway())
        files = sorted(a["file"].split("/")[-1] for a in self.ledger()["attempts"])
        self.assertEqual(len(files), 3)
        self.assertEqual(len(set(files)), 3)

    def test_the_workbook_keeps_user_edits_when_takes_are_added(self):
        import openpyxl
        self.prepare(2)
        ex.refresh(self.exp)                                             # a workbook exists
        path = self.exp / "experiment.xlsx"
        wb = openpyxl.load_workbook(path)
        ws = wb["Experiment"]
        head = [c.value for c in ws[1]]
        ws.cell(row=3, column=head.index("Notes") + 1).value = "keep me"
        wb.save(path)
        ex.add_takes(self.exp, ["r001"], 1)
        r2 = next(r for r in self.ledger()["rows"] if r["id"] == "r002")
        self.assertEqual(r2["notes"], "keep me")

    def test_a_copied_row_in_the_workbook_with_another_take_is_a_new_row(self):
        import openpyxl
        self.prepare(1)
        ex.refresh(self.exp)
        path = self.exp / "experiment.xlsx"
        wb = openpyxl.load_workbook(path)
        ws = wb["Experiment"]
        head = [c.value for c in ws[1]]
        row = [c.value for c in ws[2]]
        ws.append(row)                                                   # Excel copy-paste: same ID, same values
        ws.cell(row=ws.max_row, column=head.index("Take") + 1).value = 2
        wb.save(path)
        res = ex.refresh(self.exp)
        self.assertTrue(any("duplicate ID r001" in w for w in res["warnings"]))
        copy_row = [r for r in self.ledger()["rows"] if r["take"] == 2][0]
        self.assertNotEqual(copy_row["id"], "r001")
        self.assertEqual(copy_row["params"], self.ledger()["rows"][0]["params"])

    def test_identical_rows_with_the_same_take_are_flagged_different_takes_are_not(self):
        self.prepare(2)
        led = self.ledger()
        led["rows"][1]["params"] = dict(led["rows"][0]["params"])
        lg.save(self.exp / "ledger.json", led)
        rep = pf.run_preflight(self.exp, advisor=ADVISOR, write=False)
        self.assertTrue(any("same values, model, reference and Take" in w for w in rep["warnings"]))
        led["rows"][1]["take"] = 2
        lg.save(self.exp / "ledger.json", led)
        rep = pf.run_preflight(self.exp, advisor=ADVISOR, write=False)
        self.assertFalse(any("same values, model, reference and Take" in w for w in rep["warnings"]))

    def test_an_unknown_row_is_refused(self):
        self.prepare(1)
        with self.assertRaises(mx.PlanError):
            ex.add_takes(self.exp, ["r999"], 1)

    def test_the_cli_adds_takes(self):
        self.prepare(1)
        r = subprocess.run([sys.executable, str(SCRIPTS / "image.py"), "add-takes", "--dir", str(self.exp),
                            "--rows", "r001", "--count", "2"], capture_output=True, text=True, encoding="utf-8",
                           env={**__import__("os").environ, "PYTHONUTF8": "1"})
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("(Take 2)", r.stdout)
        self.assertIn("(Take 3)", r.stdout)


class ReviewFixTests(Base):
    """Regressions for the review of the takes, ledger and preflight code (2026-10-08)."""
    prepare = tg.GenerateTests.prepare
    run_gen = tg.GenerateTests.run_gen

    def remove_and_restore(self, rid):
        led = self.ledger()
        next(r for r in led["rows"] if r["id"] == rid)["status"] = "Removed"
        lg.save(self.exp / "ledger.json", led)
        led = self.ledger()
        next(r for r in led["rows"] if r["id"] == rid)["status"] = "Draft"       # restored by a merge
        lg.save(self.exp / "ledger.json", led)

    def test_a_restored_row_with_an_image_is_not_billed_again(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway())
        self.remove_and_restore("r001")
        rep = pf.run_preflight(self.exp, advisor=ADVISOR, write=False)
        self.assertEqual(rep["ready"], 0)
        self.assertTrue(any("already has its image" in w for w in rep["warnings"]))

    def test_a_restored_row_with_an_unknown_outcome_is_not_resubmitted(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway(lambda n, b: (0, {}, "TimeoutError")), concurrency=1)
        self.remove_and_restore("r001")
        rep = pf.run_preflight(self.exp, advisor=ADVISOR, write=False)
        self.assertEqual(rep["ready"], 0)
        self.assertTrue(any("unknown" in w.lower() for w in rep["warnings"]), rep["warnings"])

    def test_a_billed_but_not_downloaded_image_is_never_regenerated_by_a_plain_preflight(self):
        fp = self.prepare(1)

        class Broken(FakeGateway):
            def download(self, url, dest):
                raise OSError("tunnel connection failed")

        self.run_gen(fp, Broken())
        rep = pf.run_preflight(self.exp, advisor=ADVISOR, write=False)
        self.assertEqual(rep["ready"], 0)
        self.assertTrue(any("recover" in w for w in rep["warnings"]), rep["warnings"])

    def test_a_failed_row_without_a_billed_image_is_still_generated_again(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway(fail_reasons={"req-1": "boom"}), concurrency=1)
        self.assertEqual(pf.run_preflight(self.exp, advisor=ADVISOR, write=False)["ready"], 1)

    def test_an_unreadable_or_empty_lock_counts_as_held_while_it_is_recent(self):
        self.build()
        lock = self.exp / lg.LOCK_NAME
        lock.write_text("", encoding="utf-8")
        with self.assertRaises(lg.LedgerBusy):
            lg.assert_unlocked(self.exp)
        old = time.time() - 3600
        os.utime(lock, (old, old))
        lg.assert_unlocked(self.exp)                                   # an old empty file is a crash: not busy

    def test_the_heartbeat_never_leaves_an_empty_lock(self):
        self.build()
        with lg.Lock(self.exp, "test") as held:
            for _ in range(20):
                held.touch()
                self.assertTrue((self.exp / lg.LOCK_NAME).read_text(encoding="utf-8").strip())

    def test_quick_and_recover_hold_the_lock_like_run(self):
        self.build()
        with lg.Lock(self.exp, "run test"):
            with self.assertRaises(lg.LedgerBusy):
                ex.build_quick("a cat", "e-commerce product photo", 1, self.exp)

    def test_rows_added_while_the_workbook_fell_back_are_not_removed_by_the_next_merge(self):
        self.prepare(2)
        ex.refresh(self.exp)                                           # the main workbook holds the current rows
        led = self.ledger()
        before = set(led["experiment"]["main_workbook_rows"])
        new = lg.add_takes(led, ["r001"], 1)[0]                        # added while the main workbook could not be written
        lg.save(self.exp / "ledger.json", led)
        res = ex.refresh(self.exp)                                     # merges the stale main workbook
        self.assertNotIn("Removed", {r["status"] for r in self.ledger()["rows"] if r["id"] == new["id"]})
        self.assertTrue(before <= set(self.ledger()["experiment"]["main_workbook_rows"]))
        self.assertIn(new["id"], self.ledger()["experiment"]["main_workbook_rows"])      # now written, so it can be removed later
        self.assertEqual(res["warnings"], [])

    def test_take_rows_added_by_the_legacy_migration_survive_an_old_workbook(self):
        import tempfile
        tmp = Path(tempfile.mkdtemp())
        led = lg.new_ledger({"brief": "b", "intent": "x"}, ["camera"], [("a",)])
        for r in led["rows"]:
            r.pop("take")
            r["qty"] = 3
        lg.write_json_atomic(tmp / "ledger.json", led)
        got = lg.load(tmp / "ledger.json")
        self.assertEqual(len(got["rows"]), 3)
        old_workbook = [{"line": 2, "id": "r001", "selected": True, "params": {"camera": "a"}}]
        lg.merge_workbook(got, old_workbook)                           # an older workbook has only the first row
        self.assertEqual([r["status"] for r in got["rows"]].count("Removed"), 0)

    def test_a_run_records_the_model_only_for_its_own_rows_and_a_cleared_cell_stays_cleared(self):
        fp = self.prepare(2)
        led = self.ledger()
        led["rows"][0]["model"] = None
        lg.save(self.exp / "ledger.json", led)
        ex.refresh(self.exp, record_models={"r002"})                   # only r002's batch finished
        rows = {r["id"]: r for r in self.ledger()["rows"]}
        self.assertIsNone(rows["r001"].get("model"))

    def test_a_copy_of_a_row_that_already_has_its_image_is_called_out(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway())
        led = self.ledger()
        twin = lg.add_row(led, led["rows"][0]["params"], selected=True, model=led["rows"][0]["model"])   # a copy in Excel: same values and model, Take 1
        lg.save(self.exp / "ledger.json", led)
        rep = pf.run_preflight(self.exp, advisor=ADVISOR, write=False)
        self.assertTrue(any(twin["id"] in w and "already has its image" in w and "same values and Take" in w for w in rep["warnings"]),
                        rep["warnings"])

    def test_take_numbers_are_per_model_and_a_bad_take_does_not_crash_add_takes(self):
        led = lg.new_ledger({"brief": "b", "intent": "x"}, ["camera"], [("a",), ("a",)])
        led["rows"][0]["model"], led["rows"][1]["model"] = "m1", "m2"
        made = lg.add_takes(led, ["r001"], 2)
        self.assertEqual([r["take"] for r in made], [2, 3])                     # m2's row does not count as a take of m1's
        led["rows"][0]["take"] = "two"
        self.assertEqual(len(lg.add_takes(led, ["r001"], 1)), 1)                 # no TypeError

    def test_add_takes_in_a_quick_folder_says_what_to_do(self):
        out = self.tmp / "q"
        ex.build_quick("a cat", "e-commerce product photo", 1, out, models=["openai/gpt-image-2.5-sunburst"])
        with self.assertRaises(mx.PlanError) as cm:
            ex.add_takes(out, ["r001"], 1)
        self.assertIn("quick", str(cm.exception))


class PlanTakesTests(unittest.TestCase):
    def test_plan_takes_is_validated(self):
        for bad in (0, 4, "2", True, 1.5):
            p = copy.deepcopy(PLAN)
            p["takes"] = bad
            self.assertTrue(any("takes must be" in e for e in mx.validate_plan(p)), bad)
        p = copy.deepcopy(PLAN)
        p["takes"] = 2
        self.assertEqual(mx.validate_plan(p), [])

    def test_the_preview_counts_images_not_configurations(self):
        p = copy.deepcopy(PLAN)
        base = ex.preview_plan(p)
        p["takes"] = 3
        s = ex.preview_plan(p)
        self.assertEqual(s["takes"], 3)
        self.assertIn(f"{s['recommended']} configurations × 3 takes = {s['recommended'] * 3} images",
                      ex.format_plan_summary(s))
        self.assertIn("one image each", ex.format_plan_summary(base))
        if base["estimate"]["known_usd"]:
            self.assertAlmostEqual(s["estimate"]["known_usd"], base["estimate"]["known_usd"] * 3, places=3)


class MigrationTests(unittest.TestCase):
    def test_an_older_ledger_with_qty_loads_as_takes(self):
        import tempfile
        tmp = Path(tempfile.mkdtemp())
        led = lg.new_ledger({"brief": "b", "intent": "x"}, ["camera"], [("a",), ("b",), ("c",)])
        for r in led["rows"]:
            r.pop("take")
            r["qty"] = 1
        led["rows"][0]["qty"] = 3                                           # asked for 3 images, nothing generated yet
        led["rows"][1]["qty"] = 2
        led["attempts"].append({"row_id": "r002", "status": "Completed", "sample": 1, "sample_id": "r002-1"})
        lg.write_json_atomic(tmp / "ledger.json", led)
        got = lg.load(tmp / "ledger.json")
        self.assertTrue(all("qty" not in r and r["take"] >= 1 for r in got["rows"]))
        self.assertEqual(sorted(r["take"] for r in got["rows"] if r["params"] == {"camera": "a"}), [1, 2, 3])
        self.assertEqual(len([r for r in got["rows"] if r["params"] == {"camera": "b"}]), 1)   # it already has an image: no extras
        self.assertEqual(len(got["rows"]), 5)


class QuickTakesTests(Base):
    def test_quick_images_are_rows_with_takes(self):
        out = self.tmp / "q"
        q = ex.build_quick("a cat", "e-commerce product photo", 2, out, models=["openai/gpt-image-2.5-sunburst"])
        led = lg.load(out / "ledger.json")
        self.assertEqual([r["take"] for r in led["rows"]], [1, 2])
        self.assertEqual(q["images"], 2)
        self.assertEqual({r["qty"] for r in q["rows"]}, {1})
        self.assertNotIn("x2", ex.format_quick(q))


if __name__ == "__main__":
    unittest.main()
