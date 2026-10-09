"""Offline tests for scripts/generate.py: a fake gateway drives the whole run, nothing is spent."""
import json
import sys
import threading
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "scripts"))
sys.path.insert(0, str(HERE))
import generate as gen  # noqa: E402
import ledger as lg  # noqa: E402
import preflight as pf  # noqa: E402
from test_preflight import Base, PLAN, PNG, NB, ADVISOR  # noqa: E402


class FakeGateway:
    """submit_script: callable(n, body) -> (status, json, err) or None for a normal accept.
    A sample completes on its `polls`-th poll."""

    def __init__(self, submit_script=None, polls=1, cost=0.039, fail_reasons=None):
        self.bodies, self.n, self.polls, self.cost = [], 0, polls, cost
        self.script, self.fail_reasons = submit_script, fail_reasons or {}
        self.lock = threading.Lock()
        self.seen = {}

    def submit(self, body):
        with self.lock:
            self.n += 1
            n = self.n
            self.bodies.append(body)
        if self.script:
            r = self.script(n, body)
            if r is not None:
                return r
        return 200, {"data": {"request_id": f"req-{n}"}}, ""

    def poll(self, rid):
        with self.lock:
            self.seen[rid] = self.seen.get(rid, 0) + 1
            k = self.seen[rid]
        if k < self.polls:
            return 200, {"data": {"status": "IN_PROGRESS", "progress": "50%"}}, ""
        if rid in self.fail_reasons:
            return 200, {"data": {"status": "FAILED", "fail_reason": self.fail_reasons[rid]}}, ""
        return 200, {"data": {"status": "COMPLETED", "cost": self.cost, "start_time": 1, "finish_time": 5,
                              "data": {"image_urls": [f"https://x/{rid}.png"]}}}, ""

    def download(self, url, dest):
        Path(dest).write_bytes(PNG)


def quiet(_):
    pass


class GenerateTests(Base):
    def prepare(self, n=4):
        """A small experiment: the first n rows ticked, the rest not (a row is one image)."""
        self.build()
        led = self.ledger()
        for i, r in enumerate(led["rows"]):
            r["selected"] = i < n
        lg.save(self.exp / "ledger.json", led)
        (self.exp / "experiment.xlsx").unlink()
        rep = pf.run_preflight(self.exp, advisor=ADVISOR)
        self.assertEqual(rep["ready"], n)
        return rep["fingerprint"]

    def run_gen(self, fp, gw, **kw):
        kw.setdefault("poll_seconds", 0.01)
        kw.setdefault("record_prices", False)
        return gen.run_batch(self.exp, fp, gateway=gw, log=quiet, **kw)

    def test_happy_path_completes_every_sample_and_records_everything(self):
        fp = self.prepare(4)
        res = self.run_gen(fp, FakeGateway())
        self.assertEqual(res["by_status"], {"Completed": 4})
        led = self.ledger()
        self.assertEqual({r["status"] for r in led["rows"] if r["selected"]}, {"Completed"})
        self.assertEqual(len(led["batches"]), 1)
        self.assertIsNotNone(led["batches"][0]["ended_at"])
        for a in led["attempts"]:
            self.assertTrue(a["request_id"].startswith("req-"))
            self.assertTrue(a["request_hash"])
            self.assertTrue((self.exp / a["file"]).exists())
        self.assertAlmostEqual(res["actual_usd"], 4 * 0.039, places=6)
        self.assertTrue((self.exp / "round-1" / "run.json").exists())

    def test_reference_is_sent_as_a_data_uri(self):
        fp = self.prepare(1)
        gw = FakeGateway()
        self.run_gen(fp, gw)
        self.assertTrue(gw.bodies[0]["image"].startswith("data:image/png;base64,"))
        self.assertEqual(gw.bodies[0]["model"], NB)
        self.assertNotIn("watermark", gw.bodies[0])

    def test_attempt_and_request_id_are_in_the_ledger_before_the_next_step(self):
        fp = self.prepare(1)
        snaps = []
        exp = self.exp

        class Spy(FakeGateway):
            def submit(self, body):
                snaps.append(("at_submit", [dict(a) for a in lg.load(exp / "ledger.json")["attempts"]]))
                return super().submit(body)

            def poll(self, rid):
                snaps.append(("at_poll", [dict(a) for a in lg.load(exp / "ledger.json")["attempts"]]))
                return super().poll(rid)

        self.run_gen(fp, Spy())
        at_submit = next(s for k, s in snaps if k == "at_submit")
        self.assertEqual(at_submit[0]["status"], "Submitting")        # written BEFORE the POST
        self.assertTrue(at_submit[0]["request_hash"])
        at_poll = next(s for k, s in snaps if k == "at_poll")
        self.assertEqual(at_poll[0]["request_id"], "req-1")           # request id first, then polling

    def test_snapshot_must_match_its_fingerprint(self):
        fp = self.prepare(2)
        p = self.exp / "snapshots" / f"{fp}.json"
        snap = json.loads(p.read_text(encoding="utf-8"))
        snap["rows"][0]["prompt"] += " tampered"
        p.write_text(json.dumps(snap), encoding="utf-8")
        gw = FakeGateway()
        with self.assertRaises(gen.RunRefused):
            self.run_gen(fp, gw)
        self.assertEqual(gw.n, 0)

    def test_changed_reference_refuses_before_any_spend(self):
        fp = self.prepare(2)
        (self.exp / "refs" / "mighty-product.png").write_bytes(PNG + b"x")
        gw = FakeGateway()
        with self.assertRaises(gen.RunRefused):
            self.run_gen(fp, gw)
        self.assertEqual(gw.n, 0)

    def test_unknown_fingerprint_refused(self):
        self.prepare(1)
        with self.assertRaises(gen.RunRefused):
            self.run_gen("deadbeef0000", FakeGateway())

    def test_edits_after_approval_cannot_change_the_run(self):
        fp = self.prepare(2)
        led = self.ledger()
        for r in led["rows"]:
            r["params"]["lighting"] = "something else"
            r["selected"] = True
        lg.save(self.exp / "ledger.json", led)
        gw = FakeGateway()
        self.run_gen(fp, gw)
        self.assertEqual(gw.n, 2)                                     # exactly the approved rows
        self.assertNotIn("something else", gw.bodies[0]["prompt"])

    def test_timeout_on_submit_is_unknown_and_never_retried(self):
        fp = self.prepare(3)
        gw = FakeGateway(lambda n, body: (0, {}, "TimeoutError: read timed out") if n == 2 else None)
        res = self.run_gen(fp, gw)
        self.assertEqual(res["by_status"], {"Completed": 2, "Unknown": 1})
        self.assertEqual(gw.n, 3)                                     # submitted once each, never again
        self.assertEqual(len(res["unknown"]), 1)
        self.assertTrue(res["unknown"][0]["request_hash"])
        self.assertIn("Unknown", {r["status"] for r in self.ledger()["rows"]})
        rep = pf.run_preflight(self.exp, advisor=ADVISOR)             # preflight does not pick it up again
        self.assertTrue(any("Unknown" in w and "never retried" in w for w in rep["warnings"]))

    def test_crash_while_submitting_becomes_unknown_on_resume(self):
        fp = self.prepare(2)
        led = self.ledger()
        snap = gen.load_snapshot(self.exp, fp)
        led["batches"].append({"no": 1, "fingerprint": fp, "approved_at": lg.now(), "models": [],
                               "estimated_usd": 0.078, "max_usd": 1, "started_at": lg.now(),
                               "ended_at": None, "actual_usd": 0})
        for i, srow in enumerate(snap["rows"]):
            led["attempts"].append({"batch": 1, "row_id": srow["id"], "sample": 1, "sample_id": srow["id"] + "-1",
                                    "model": srow["model"], "mode": "reference", "size": None, "est_usd": 0.039,
                                    "status": "Submitting" if i == 0 else "Pending", "request_id": None,
                                    "request_hash": "abc", "submitted_at": lg.now()})
        lg.save(self.exp / "ledger.json", led)
        gw = FakeGateway()
        res = self.run_gen(fp, gw)
        self.assertEqual(res["by_status"], {"Unknown": 1, "Completed": 1})
        self.assertEqual(gw.n, 1)                                     # the interrupted one was NOT resubmitted

    def test_resume_polls_running_samples_without_resubmitting(self):
        fp = self.prepare(2)
        res = self.run_gen(fp, FakeGateway(polls=10_000), poll_timeout=0.05)   # too short: they stay Running
        self.assertEqual(sorted(res["unfinished"]), ["r001", "r002"])
        self.assertIsNone(self.ledger()["batches"][0]["ended_at"])
        gw2 = FakeGateway()
        res2 = self.run_gen(fp, gw2)
        self.assertEqual(gw2.n, 0)                                    # no new submissions
        self.assertEqual(res2["by_status"], {"Completed": 2})
        self.assertIsNotNone(self.ledger()["batches"][0]["ended_at"])

    def test_same_snapshot_cannot_be_rerun_without_again(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway())
        with self.assertRaises(gen.RunRefused):
            self.run_gen(fp, FakeGateway())
        gw = FakeGateway()
        self.run_gen(fp, gw, again=True)
        self.assertEqual(gw.n, 1)
        self.assertEqual(len(self.ledger()["batches"]), 2)

    def test_auth_rejection_stops_everything_after_it(self):
        fp = self.prepare(6)
        gw = FakeGateway(lambda n, body: (401, {"error": {"code": "invalid_api_key", "message": "bad key"}}, "")
                         if n >= 2 else None)
        res = self.run_gen(fp, gw, concurrency=1)
        self.assertLessEqual(gw.n, 2)                                 # the breaker stopped further submissions
        self.assertIn("rejected", res["stopped"])
        self.assertEqual(res["by_status"].get("Pending"), 4)
        self.assertNotIn("Queued", {r["status"] for r in self.ledger()["rows"]})

    def test_budget_guard_counts_in_flight_samples(self):
        fp = self.prepare(6)
        gw = FakeGateway(polls=3)
        res = self.run_gen(fp, gw, concurrency=6, max_usd=0.10)       # 0.039 each: room for 2 only
        self.assertEqual(gw.n, 2)
        self.assertIn("budget guard", res["stopped"])
        self.assertEqual(res["by_status"]["Completed"], 2)

    def test_budget_guard_also_stops_rows_with_no_price_basis(self):
        fp = self.prepare(4)
        snap_p = self.exp / "snapshots" / f"{fp}.json"
        snap = json.loads(snap_p.read_text(encoding="utf-8"))
        for r in snap["rows"]:
            r["est_usd"] = None                                       # unverified prices: no estimate
        snap["unverified_rows"] = [r["id"] for r in snap["rows"]]
        snap["fingerprint"] = pf.snapshot_fingerprint(snap)
        fp2 = snap["fingerprint"]
        (self.exp / "snapshots" / f"{fp2}.json").write_text(json.dumps(snap), encoding="utf-8")
        gw = FakeGateway(cost=0.05)
        res = self.run_gen(fp2, gw, concurrency=1, max_usd=0.12)      # the 1st (unpriced) bills 0.05, which is then the
        self.assertEqual(gw.n, 2)                                     # price of the rest: 2 fit, the 3rd would not
        self.assertIn("budget guard", res["stopped"])

    def test_default_budget_is_derived_and_unverified_requires_explicit(self):
        fp = self.prepare(2)
        self.assertIsNone(self.run_gen(fp, FakeGateway())["stopped"])
        snap_p = self.exp / "snapshots" / f"{fp}.json"
        snap = json.loads(snap_p.read_text(encoding="utf-8"))
        snap["unverified_rows"] = ["r001"]                            # tampering without re-hashing is refused...
        snap_p.write_text(json.dumps(snap), encoding="utf-8")
        with self.assertRaises(gen.RunRefused):
            self.run_gen(fp, FakeGateway(), again=True)
        snap["fingerprint"] = pf.snapshot_fingerprint(snap)           # ...a legitimately re-approved one needs --max-usd
        (self.exp / "snapshots" / f"{snap['fingerprint']}.json").write_text(json.dumps(snap), encoding="utf-8")
        with self.assertRaises(gen.RunRefused) as cm:
            self.run_gen(snap["fingerprint"], FakeGateway())
        self.assertIn("--max-usd", str(cm.exception))

    def test_moderation_failure_is_blocked_and_other_failures_are_failed(self):
        fp = self.prepare(3)
        gw = FakeGateway(fail_reasons={"req-1": "Blocked by safety policy", "req-2": "internal error"})
        res = self.run_gen(fp, gw, concurrency=1)
        self.assertEqual(res["by_status"], {"Blocked": 1, "Failed": 1, "Completed": 1})
        self.assertEqual(sorted(r["status"] for r in self.ledger()["rows"] if r["selected"]),
                         ["Blocked", "Completed", "Failed"])

    def test_takes_are_separate_rows_each_one_image(self):
        self.prepare(1)
        led = self.ledger()
        lg.add_takes(led, ["r001"], 2)
        lg.save(self.exp / "ledger.json", led)
        fp = pf.run_preflight(self.exp, advisor=ADVISOR)["fingerprint"]
        res = self.run_gen(fp, FakeGateway())
        self.assertEqual(res["samples"], 3)
        led = self.ledger()
        picked = [r for r in led["rows"] if r["selected"]]
        self.assertEqual(sorted(r["take"] for r in picked), [1, 2, 3])
        self.assertEqual(len({str(r["params"]) for r in picked}), 1)              # identical values
        self.assertEqual(sorted(a["sample_id"] for a in led["attempts"]), sorted(r["id"] for r in picked))
        for a in led["attempts"]:
            self.assertTrue(a["file"].endswith(a["row_id"] + ".png"))             # r001.png, not r001-1.png

    def test_a_failed_sample_fails_its_row(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway(fail_reasons={"req-1": "boom"}), concurrency=1)
        self.assertEqual(self.ledger()["rows"][0]["status"], "Failed")

    def test_refresh_after_run_adds_image_links_and_costs(self):
        import experiment as ex
        import openpyxl
        fp = self.prepare(2)
        self.run_gen(fp, FakeGateway())
        ex.refresh(self.exp)
        ws = openpyxl.load_workbook(self.exp / "experiment.xlsx")["Experiment"]
        head = [c.value for c in ws[1]]
        row2 = [c.value for c in ws[2]]
        self.assertEqual(row2[head.index("Status")], "Completed")
        self.assertAlmostEqual(row2[head.index("Cost")], 0.039)
        self.assertEqual(row2[head.index("File")], "r001.png")

    def test_model_used_is_recorded_in_the_model_column(self):
        import experiment as ex
        import openpyxl
        fp = self.prepare(2)
        led = self.ledger()
        self.assertTrue(all(not r.get("model") for r in led["rows"]))      # the Advisor picks, the cell is blank
        self.run_gen(fp, FakeGateway())
        led = self.ledger()
        used = {a["row_id"]: a["model"] for a in led["attempts"]}
        for r in led["rows"]:
            if r["selected"]:
                self.assertEqual(r["model"], used[r["id"]])
            else:
                self.assertFalse(r.get("model"))                            # not generated: still blank
        ex.refresh(self.exp)
        ws = openpyxl.load_workbook(self.exp / "experiment.xlsx")["Experiment"]
        head = [c.value for c in ws[1]]
        self.assertEqual(ws[2][head.index("Model")].value, used["r001"])

    def test_model_survives_a_workbook_written_before_the_run(self):
        """Regression: after a run the old workbook (blank Model) was merged back and blanked the recorded model."""
        import experiment as ex
        import openpyxl
        fp = self.prepare(2)
        ex.refresh(self.exp)                                               # the workbook the user has: Model blank
        fp = pf.run_preflight(self.exp, advisor=ADVISOR)["fingerprint"]
        self.run_gen(fp, FakeGateway())
        ex.after_run(self.exp)
        ws = openpyxl.load_workbook(self.exp / "experiment.xlsx")["Experiment"]
        head = [c.value for c in ws[1]]
        self.assertTrue(ws[2][head.index("Model")].value)
        self.assertTrue(self.ledger()["rows"][0]["model"])

    def test_record_models_flag_fills_an_older_experiment(self):
        import experiment as ex
        fp = self.prepare(2)
        self.run_gen(fp, FakeGateway())
        led = self.ledger()
        for r in led["rows"]:
            r["model"] = None
        lg.save(self.exp / "ledger.json", led)
        ex.refresh(self.exp)
        self.assertFalse(self.ledger()["rows"][0].get("model"))           # a plain refresh never re-fills
        ex.refresh(self.exp, record_models=True)
        self.assertTrue(self.ledger()["rows"][0]["model"])

    def test_a_models_the_user_chose_is_not_overwritten_and_a_cleared_one_stays_cleared_on_refresh(self):
        import experiment as ex
        fp = self.prepare(2)
        led = self.ledger()
        led["rows"][1]["model"] = NB
        lg.save(self.exp / "ledger.json", led)
        fp = pf.run_preflight(self.exp, advisor=ADVISOR)["fingerprint"]
        self.run_gen(fp, FakeGateway())
        led = self.ledger()
        self.assertEqual(led["rows"][1]["model"], NB)
        led["rows"][0]["model"] = None                                      # the user clears it afterwards
        lg.save(self.exp / "ledger.json", led)
        ex.refresh(self.exp)
        self.assertFalse(self.ledger()["rows"][0].get("model"))

    def test_a_row_with_no_completed_sample_gets_no_model(self):
        fp = self.prepare(1)
        self.run_gen(fp, FakeGateway(fail_reasons={"req-1": "boom"}))
        self.assertFalse(self.ledger()["rows"][0].get("model"))

    def test_ticked_completed_rows_are_skipped_and_pointed_to_add_takes(self):
        fp = self.prepare(2)
        self.run_gen(fp, FakeGateway())
        rep = pf.run_preflight(self.exp, advisor=ADVISOR)
        self.assertEqual(rep["ready"], 0)                                          # nothing is billed twice
        self.assertTrue(any("already has its image" in w and "add-takes" in w for w in rep["warnings"]))
        # one line for all of them, however many rows: the text no longer names each row's file
        lines = [w for w in pf.group_messages(rep["warnings"]) if "already has its image" in w]
        self.assertEqual(len(lines), 1)
        self.assertIn("r001, r002", lines[0])


if __name__ == "__main__":
    unittest.main()


class SpendingLimitReport(unittest.TestCase):
    BASE = {"batch": 1, "out_dir": "round-1", "samples": 4, "by_status": {"Completed": 4}, "estimated_usd_known": 0.05, "actual_usd": 0.06,
            "by_model_usd": {"m": 0.06}, "stopped": None, "unfinished": [], "unknown": [], "failed": []}

    def test_the_result_says_the_limit_is_not_a_hard_cap_and_reports_an_overshoot(self):
        text = gen.format_result({**self.BASE, "max_usd": 0.05})
        self.assertIn("limits NEW submissions; it is not a hard cap", text)
        self.assertIn("Actual cost is $0.0100 over it", text)

    def test_no_overshoot_no_overshoot_line_and_no_limit_no_limit_line(self):
        self.assertNotIn("over it", gen.format_result({**self.BASE, "max_usd": 0.10}))
        self.assertNotIn("Spending limit", gen.format_result({**self.BASE, "max_usd": None}))
